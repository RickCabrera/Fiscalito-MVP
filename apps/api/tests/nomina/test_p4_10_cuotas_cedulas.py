"""
P4-10: cuotas mensuales (EMA) y bimestrales (EBA) desde el historial de SBC y las nóminas
cerradas; importación de la cédula y conciliación por trabajador; reporte exportable.

El caso a mano: SBC $420.00 (3.58 UMA → CEAV 3.51-4.00 UMA = 6.613 %), prima RT 0.54355 %,
UMA 2026 $117.31 (3 UMA = $351.93), zona general, redondeo por concepto (§D2, ROUND_HALF_UP).
"""

from __future__ import annotations

import io
import json
from datetime import date
from decimal import Decimal as D

import pytest
from openpyxl import Workbook, load_workbook

from app.constants import ZonaSalarioMinimo
from app.modelos.nomina import NominaPeriodo
from app.nomina_engine.cuotas import DiasDelPeriodo, cuotas_empleado
from app.nomina_engine.cuotas_mensuales import TramoSBC, cuotas_del_mes
from app.nomina_engine.integracion import clamp_sbc
from tests.nomina.test_p4_09_movimientos_imss import _emp, _sembrar

BASE = "/api/v1/nomina/clientes/imss"
PRIMA = D("0.0054355")
GENERAL = ZonaSalarioMinimo.GENERAL

# Junio 2026, 30 días, sin ausencias — a mano.
EMA_JUNIO = {
    "eym_cuota_fija": ("717.94", "0.00"),       # 0.204 × 117.31 × 30 = 717.9372
    "eym_excedente": ("22.46", "8.17"),         # (420 − 351.93) × 30 × 1.1 % / 0.4 %
    "eym_prestaciones_dinero": ("88.20", "31.50"),
    "eym_gastos_medicos_pensionados": ("132.30", "47.25"),
    "invalidez_vida": ("220.50", "78.75"),
    "guarderias": ("126.00", "0.00"),
    "riesgos_trabajo": ("68.49", "0.00"),       # 420 × 0.0054355 × 30 = 68.4873
}
TOTAL_EMA_JUNIO = D("1541.56")
# Bimestre 3 (mayo 31 + junio 30): cada mes se redondea por concepto y luego se suma.
EBA_B3 = {
    "retiro": ("512.40", "0.00"),               # 260.40 + 252.00
    "ceav": ("1694.25", "288.23"),              # 861.01 + 833.24 / 146.48 + 141.75
    "infonavit": ("1281.00", "0.00"),           # 651.00 + 630.00
}
TOTAL_EBA_B3 = D("3775.88")


@pytest.fixture
def http(api):
    c = api("uid-cuotas")
    _sembrar(c)
    return c


class TestMotor:
    def test_mes_completo_cuadra_a_mano(self):
        c = cuotas_del_mes((TramoSBC(date(2026, 3, 2), D("420.00")),), 2026, 6, GENERAL, PRIMA)
        ramos = {r.clave: (str(r.patron), str(r.obrero)) for r in c.ramos}
        for clave, esperado in EMA_JUNIO.items():
            assert ramos[clave] == esperado, clave
        assert c.dias == 30 and c.dias_cotizados == 30

    def test_un_tramo_es_exactamente_el_motor_existente(self):
        c = cuotas_del_mes((TramoSBC(date(2026, 3, 2), D("420.00")),), 2026, 6, GENERAL, PRIMA,
                           dias_ausentismo=2, dias_incapacidad=1)
        motor = cuotas_empleado(clamp_sbc(D("420.00"), date(2026, 6, 1), GENERAL),
                                DiasDelPeriodo(30, 2, 1), date(2026, 6, 1), GENERAL, PRIMA)
        assert {r.clave: r.patron + r.obrero for r in c.ramos} == {
            r.clave: r.patron + r.obrero for r in motor.ramos}

    def test_cambio_de_sbc_a_medio_mes_parte_en_tramos(self):
        tramos = (TramoSBC(date(2026, 3, 2), D("420.00")), TramoSBC(date(2026, 6, 16), D("500.00")))
        c = cuotas_del_mes(tramos, 2026, 6, GENERAL, PRIMA)
        assert [(s.desde.day, s.dias, s.sbc) for s in c.segmentos] == [
            (1, 15, D("420.00")), (16, 15, D("500.00"))]
        iv = next(r for r in c.ramos if r.clave == "invalidez_vida")
        assert iv.patron == D("110.25") + D("131.25")  # 420×1.75 %×15 + 500×1.75 %×15

    def test_alta_a_medio_mes_cuenta_dias_naturales_vigentes(self):
        c = cuotas_del_mes((TramoSBC(date(2026, 6, 10), D("420.00")),), 2026, 6, GENERAL, PRIMA,
                           alta=date(2026, 6, 10))
        assert c.dias == 21
        assert cuotas_del_mes((), 2026, 5, GENERAL, PRIMA, alta=date(2026, 6, 10)) is None

    def test_tope_de_7_dias_de_ausentismo_por_mes(self):
        c = cuotas_del_mes((TramoSBC(date(2026, 3, 2), D("420.00")),), 2026, 6, GENERAL, PRIMA,
                           dias_ausentismo=10)
        assert c.dias_ausentismo == 7 and c.avisos
        sin = cuotas_del_mes((TramoSBC(date(2026, 3, 2), D("420.00")),), 2026, 6, GENERAL, PRIMA,
                             dias_ausentismo=10, tope_ausentismo=None)
        assert sin.dias_ausentismo == 10


class TestRutas:
    def test_ema_y_eba_cuadran_a_mano(self, http):
        r = http.get(f"{BASE}/cuotas-imss", params={"ejercicio": 2026, "mes": 6})
        assert r.status_code == 200, r.text
        (t,) = r.json()["trabajadores"]
        assert {k: (v["patron"], v["obrero"]) for k, v in t["ramos"].items()} == EMA_JUNIO
        assert D(t["total"]) == TOTAL_EMA_JUNIO and t["dias"] == 30
        r = http.get(f"{BASE}/cuotas-imss", params={"ejercicio": 2026, "bimestre": 3})
        (t,) = r.json()["trabajadores"]
        assert {k: (v["patron"], v["obrero"]) for k, v in t["ramos"].items()} == EBA_B3
        assert D(t["total"]) == TOTAL_EBA_B3 and t["dias"] == 61

    def test_ausencias_salen_de_las_nominas_cerradas(self, http):
        r = http.post("/api/v1/nomina/calcular-periodo", json={
            "cliente": "imss", "periodo": {"inicio": "2026-06-01", "fin": "2026-06-15"},
            "incidencias": [{"empleado_no": "E1", "dias_periodo": 15, "faltas": 2,
                             "dias_ausentismo": 2}]})
        assert r.status_code == 200, r.text
        pid = r.json()["periodo_id"]
        # Mientras no se cierre, no cuenta.
        t = http.get(f"{BASE}/cuotas-imss", params={"ejercicio": 2026, "mes": 6}).json()
        assert t["trabajadores"][0]["dias_ausentismo"] == 0
        http.post(f"{BASE}/periodos/{pid}/cerrar")
        t = http.get(f"{BASE}/cuotas-imss", params={"ejercicio": 2026, "mes": 6}).json()
        fila = t["trabajadores"][0]
        assert fila["dias_ausentismo"] == 2 and fila["dias_cotizados"] == 28
        assert fila["ramos"]["invalidez_vida"]["patron"] == "205.80"   # 420 × 1.75 % × 28
        assert fila["ramos"]["eym_cuota_fija"]["patron"] == "717.94"   # EyM no se reduce (§D3)

    def test_csv_exportable(self, http):
        r = http.get(f"{BASE}/cuotas-imss", params={"ejercicio": 2026, "mes": 6, "formato": "csv"})
        assert r.status_code == 200 and r.headers["content-type"].startswith("text/csv")
        assert "12345678901" in r.content.decode("utf-8-sig")

    def test_sin_mes_ni_bimestre_es_422(self, http):
        assert http.get(f"{BASE}/cuotas-imss", params={"ejercicio": 2026}).status_code == 422


def _csv(filas: list[list[str]]) -> bytes:
    return "\n".join(",".join(f) for f in filas).encode()


class TestCedula:
    def _otro(self, http):
        http.put("/api/v1/cartera/clientes/imss/empleados/E2", json=_emp(
            "E2", nss="98765432109", salario_diario_integrado="420.00"))

    def test_concilia_y_marca_diferencias(self, http):
        self._otro(http)
        cuerpo = _csv([
            ["NSS", "NOMBRE", "DIAS", "SBC", "INVALIDEZ_VIDA_PATRONAL", "TOTAL"],
            ["12345678901", "PENA ALVAREZ JOSE", "30", "420.00", "220.50", "1541.56"],
            ["98765432109", "OTRO", "28", "420.00", "220.50", "1600.00"],
            ["11111111111", "FANTASMA", "30", "300.00", "", "900.00"],
        ])
        r = http.post(f"{BASE}/cedulas", data={"tipo": "EMA", "ejercicio": "2026", "periodo": "6"},
                      files={"archivo": ("ema.csv", cuerpo, "text/csv")})
        assert r.status_code == 200, r.text
        conc = r.json()["cedula"]["conciliacion"]
        por_nss = {t["nss"]: t for t in conc["trabajadores"]}
        assert por_nss["12345678901"]["cuadra"] is True
        difs = {d["campo"]: d for d in por_nss["98765432109"]["diferencias"]}
        assert set(difs) == {"dias", "importe"}
        assert difs["importe"]["diferencia"] == "58.44"
        assert por_nss["11111111111"]["diferencias"][0]["campo"] == "trabajador"
        assert conc["con_diferencia"] == 2 and r.json()["cedula"]["diferencias"] == 2

        cid = r.json()["cedula"]["id"]
        csv = http.get(f"{BASE}/cedulas/{cid}/reporte").content.decode("utf-8-sig")
        assert "98765432109" in csv and "importe" in csv
        xlsx = http.get(f"{BASE}/cedulas/{cid}/reporte", params={"formato": "xlsx"})
        hoja = load_workbook(io.BytesIO(xlsx.content)).active
        assert hoja.cell(1, 1).value == "NSS"
        assert [c["id"] for c in http.get(f"{BASE}/cedulas").json()["cedulas"]] == [cid]

    def test_eba_xlsx_con_columnas_propias(self, http):
        libro = Workbook()
        hoja = libro.active
        hoja.append(["Reporte EBA"])
        hoja.append(["Número de Seguridad Social", "Días", "Retiro", "CEAV Patrón",
                     "Amortización", "Total"])
        hoja.append(["12345678901", 61, 512.40, 1694.25, 100.00, 3875.88])
        buf = io.BytesIO()
        libro.save(buf)
        cfg = {"fila_encabezado": 2, "columnas": {
            "nss": "Número de Seguridad Social", "dias": "Días", "retiro": "Retiro",
            "ceav_patronal": "CEAV Patrón", "amortizacion": "Amortización", "total": "Total"}}
        r = http.post(f"{BASE}/cedulas", data={
            "tipo": "EBA", "ejercicio": "2026", "periodo": "3", "configuracion": json.dumps(cfg)},
            files={"archivo": ("eba.xlsx", buf.getvalue(), "application/octet-stream")})
        assert r.status_code == 200, r.text
        (t,) = r.json()["cedula"]["conciliacion"]["trabajadores"]
        assert t["cuadra"] is True, t  # total − amortización = 3,775.88

    def test_txt_de_posiciones_fijas(self, http):
        linea = "12345678901" + "30" + "0000154156"
        cfg = {"formato": "txt", "posiciones": {"nss": [1, 11], "dias": [12, 13],
                                                 "total": [14, 23]}, "decimales_implicitos": 2}
        r = http.post(f"{BASE}/cedulas", data={
            "tipo": "EMA", "ejercicio": "2026", "periodo": "6", "configuracion": json.dumps(cfg)},
            files={"archivo": ("ema.txt", linea.encode(), "text/plain")})
        assert r.status_code == 200, r.text
        assert r.json()["cedula"]["conciliacion"]["trabajadores"][0]["cuadra"] is True

    def test_sin_columna_nss_explica(self, http):
        r = http.post(f"{BASE}/cedulas", data={"tipo": "EMA", "ejercicio": "2026", "periodo": "6"},
                      files={"archivo": ("x.csv", b"A,B\n1,2\n", "text/csv")})
        assert r.status_code == 422 and "NSS" in r.json()["error"]

    def test_plantilla(self, http):
        r = http.get(f"{BASE}/cedulas/plantilla", params={"tipo": "EBA"})
        assert r.content.decode().startswith("nss,nombre,dias,sbc,retiro")

    def test_multi_tenant(self, http, api):
        otro = api("uid-otro-cuotas")
        r = otro.get(f"{BASE}/cuotas-imss", params={"ejercicio": 2026, "mes": 6})
        assert r.status_code == 404
        r = http.post(f"{BASE}/cedulas", data={"tipo": "EMA", "ejercicio": "2026", "periodo": "6"},
                      files={"archivo": ("e.csv", _csv([["NSS"], ["12345678901"]]), "text/csv")})
        cid = r.json()["cedula"]["id"]
        assert otro.get(f"{BASE}/cedulas/{cid}").status_code == 404


class TestBimestreVariable:
    def test_promedio_del_bimestre_genera_07(self, http, sesionador):
        http.put("/api/v1/cartera/clientes/imss/empleados/E1", json=_emp(tipo_salario="2"))
        r = http.post("/api/v1/nomina/calcular-periodo", json={
            "cliente": "imss", "periodo": {"inicio": "2026-04-01", "fin": "2026-04-15"},
            "incidencias": [{"empleado_no": "E1", "dias_periodo": 15, "faltas": 0,
                             "dias_ausentismo": 0}]})
        assert r.status_code == 200, r.text
        pid = r.json()["periodo_id"]
        http.post(f"{BASE}/periodos/{pid}/cerrar")
        # El integrable del periodo lo fija P4-05 por recibo; aquí se siembra el dato guardado.
        db = sesionador()
        p = db.get(NominaPeriodo, pid)
        res = dict(p.resultado)
        res["recibos"] = [{**res["recibos"][0], "sbc_variable_integrable": "1500.00"}]
        p.resultado = res
        db.commit()
        db.close()
        r = http.post(f"{BASE}/movimientos-imss/bimestre-variable",
                      json={"ejercicio": 2026, "bimestre": 2, "aplicar": True})
        assert r.status_code == 200, r.text
        (t,) = r.json()["trabajadores"]
        assert t["parte_variable"] == "100.00" and t["sbc_nuevo"] == "520.00"  # 1500 / 15
        assert t["fecha_efecto"] == "2026-05-01" and t["movimiento_id"]
        mov = next(m for m in http.get(f"{BASE}/movimientos-imss").json()["movimientos"]
                   if m["tipo"] == "07")
        assert (mov["sbc"], mov["tipo_salario"], mov["fecha"]) == ("520.00", "2", "2026-05-01")
        # y las cuotas de mayo ya usan el SBC declarado
        t = http.get(f"{BASE}/cuotas-imss", params={"ejercicio": 2026, "mes": 5}).json()
        assert t["trabajadores"][0]["sbc"] == "520.00"
