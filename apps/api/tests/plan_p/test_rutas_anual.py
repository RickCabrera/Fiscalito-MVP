"""
P3-01 / P3-06: rutas de personas morales y de la anual contra la bóveda y la contabilidad.

"Listo cuando" medidos aquí:
- P3-01: un caso con fuente (el pago provisional publicado, ver `test_personas_morales.py`)
  cuadra al peso también por la ruta, sale en el resumen de impuestos y genera sus papeles.
- P3-06: el Excel de la anual abre y sus totales cuadran con la pantalla (PF y PM).
- Calendario: la anual de una persona moral vence en marzo.
- Multi-tenant: el despacho B no ve ni calcula lo del A (404).
"""

from __future__ import annotations

import io
from datetime import date
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from app.contabilidad import catalogo, polizas
from app.contabilidad.polizas import MovimientoEntrada
from app.modelos.nucleo import Cliente
from app.tenancy import asegurar_usuario
from tests.plan_p.cfdi_sintetico import RFC_CLIENTE, RFC_COMPRADOR, RFC_PROVEEDOR, cfdi, uuid_n

D = Decimal
RFC_PM = "PMS010101AB1"
IMP = "/api/v1/impuestos/clientes"


@pytest.fixture
def despachos(sesionador, tmp_path, monkeypatch):
    monkeypatch.setenv("FISCALITO_ALMACEN_DIR", str(tmp_path / "almacen"))
    monkeypatch.delenv("FISCALITO_ALMACENAMIENTO", raising=False)
    db = sesionador()
    asegurar_usuario(db, "uid-a")
    asegurar_usuario(db, "uid-b")
    db.add_all([
        Cliente(despacho_id="uid-a", clave="moral", nombre="Moral SA", rfc=RFC_PM,
                datos={"regimen": "601"}),
        Cliente(despacho_id="uid-a", clave="resico", nombre="Resico SA", rfc=RFC_PM,
                datos={"regimen": "626"}),
        Cliente(despacho_id="uid-a", clave="fisica", nombre="Taller", rfc=RFC_CLIENTE,
                datos={"regimen": "612"}),
        Cliente(despacho_id="uid-b", clave="ajeno", nombre="Ajeno", rfc=RFC_PM,
                datos={"regimen": "601"}),
    ])
    db.commit()
    db.close()


def _subir(a, clave, archivos):
    r = a.post(f"/api/v1/boveda/clientes/{clave}/cfdi",
               files=[("archivos", (f"{i}.xml", x, "text/xml")) for i, x in enumerate(archivos)])
    assert r.status_code == 200, r.text
    assert not r.json()["rechazados"], r.json()["rechazados"]


def _venta(n, mes, subtotal, emisor=RFC_PM):
    iva = (D(subtotal) * D("0.16")).quantize(D("0.01"))
    return cfdi(uuid_n(n), emisor=emisor, receptor=RFC_COMPRADOR,
                fecha=f"2026-{mes:02d}-10T10:00:00", subtotal=subtotal,
                total=f"{D(subtotal) + iva}", iva=f"{iva}")


def _compra(n, mes, subtotal, receptor=RFC_PM):
    iva = (D(subtotal) * D("0.16")).quantize(D("0.01"))
    return cfdi(uuid_n(n, "BBBBBBBB"), emisor=RFC_PROVEEDOR, receptor=receptor,
                fecha=f"2026-{mes:02d}-12T10:00:00", subtotal=subtotal,
                total=f"{D(subtotal) + iva}", iva=f"{iva}")


def _ventas_del_ejemplo(a, clave="moral"):
    # Ejemplo publicado (siemprealdia.co): 1.1 M, 1.25 M y 1.45 M → 3.8 M a marzo.
    _subir(a, clave, [_venta(1, 1, "1100000.00"), _venta(2, 2, "1250000.00"),
                      _venta(3, 3, "1450000.00")])


DATOS_EJEMPLO = {
    "coeficiente_utilidad": "0.1845",
    "perdidas": [{"ejercicio": 2024, "remanente": "150000", "actualizada_a": "2025-12"}],
    "retenciones_bancarias": {"3": "5000"},
    "pagos_efectuados": {"1": "40000", "2": "55000"},
}


def test_pago_provisional_pm_desde_la_boveda_cuadra_con_el_ejemplo(api, despachos):
    a = api("uid-a")
    _ventas_del_ejemplo(a)
    r = a.put(f"{IMP}/moral/pm/datos/2026", json=DATOS_EJEMPLO)
    assert r.status_code == 200, r.text
    assert a.get(f"{IMP}/moral/pm/datos/2026").json()["datos"]["coeficiente_utilidad"] == "0.1845"
    r = a.post(f"{IMP}/moral/pm/pagos-provisionales", json={"anio": 2026, "mes": 3})
    assert r.status_code == 200, r.text
    d = r.json()
    assert d["tipo"] == "pm" and d["fuente"] == "boveda" and d["coeficiente"] == "0.1845"
    assert d["origen_coeficiente"] == "capturado" and d["fecha_limite"] == "2026-04-17"
    marzo = d["pagos"][-1]
    assert marzo["ingresos_acumulados"] == "3800000.00"
    assert marzo["utilidad_estimada"] == "701100.00" and marzo["isr_causado"] == "165330.00"
    assert marzo["pago_determinado"] == "65330.00"

    # El resumen de impuestos ya no sale en cero: trae el pago provisional del mes.
    res = a.post(f"{IMP}/moral/resumen",
                 json={"anio": 2026, "mes": 3, "fecha_pago": "2026-04-17"}).json()
    fila = next(f for f in res["renglones"] if f["clave"] == "isr_pm")
    assert fila["aplica"] and fila["importe_sistema"] == "65330.00"


def test_notas_de_credito_restan_y_cancelados_no_cuentan(api, despachos, sesionador):
    a = api("uid-a")
    nota = cfdi(uuid_n(9), tipo="E", emisor=RFC_PM, receptor=RFC_COMPRADOR,
                fecha="2026-01-20T10:00:00", subtotal="100000.00", total="116000.00",
                iva="16000.00")
    _subir(a, "moral", [_venta(1, 1, "1100000.00"), nota, _venta(2, 1, "5000.00")])
    from app.modelos.cfdi import Comprobante

    db = sesionador()
    db.query(Comprobante).filter_by(uuid=uuid_n(2)).one().estatus_sat = "cancelado"
    db.commit()
    db.close()
    a.put(f"{IMP}/moral/pm/datos/2026", json={"coeficiente_utilidad": "0.10"})
    ene = a.post(f"{IMP}/moral/pm/pagos-provisionales",
                 json={"anio": 2026, "mes": 1}).json()["pagos"][0]
    assert ene["ingresos_mes"] == "1000000.00" and ene["pago_determinado"] == "30000.00"


def test_sin_coeficiente_lo_calcula_de_la_anual_anterior_capturada(api, despachos):
    a = api("uid-a")
    _ventas_del_ejemplo(a)
    a.put(f"{IMP}/moral/pm/datos/2026", json={"utilidad_fiscal_base": "2700000",
                                             "ingresos_nominales_base": "12150000"})
    d = a.post(f"{IMP}/moral/pm/pagos-provisionales", json={"anio": 2026, "mes": 1}).json()
    assert d["coeficiente"] == "0.2222" and "capturados" in d["origen_coeficiente"]
    # Sin nada capturado ni ejercicios anteriores: coeficiente cero y la nota lo dice.
    a.put(f"{IMP}/moral/pm/datos/2026", json={})
    d = a.post(f"{IMP}/moral/pm/pagos-provisionales", json={"anio": 2026, "mes": 1}).json()
    assert d["coeficiente"] == "0.0000" and any("primer ejercicio" in n for n in d["notas"])


def test_anual_pm_y_papeles_cuadran_con_la_pantalla(api, despachos):
    a = api("uid-a")
    _ventas_del_ejemplo(a)
    _subir(a, "moral", [_compra(1, 2, "2000000.00")])
    a.put(f"{IMP}/moral/pm/datos/2026", json={
        **DATOS_EJEMPLO, "ptu_pagada": "80000",
        "saldos_creditos": ["1000000"] * 12, "saldos_deudas": ["1250000"] * 12,
        "deduccion_inversiones": "10000"})
    r = a.post(f"{IMP}/moral/anual", json={"anio": 2026})
    assert r.status_code == 200, r.text
    d = r.json()
    assert d["tipo"] == "pm" and d["fuente"] == "boveda" and len(d["meses"]) == 12
    ren = {x["clave"]: D(x["importe"]) for x in d["renglones"]}
    assert ren["ingresos"] == D("3800000.00")
    assert ren["deducciones"] == D("2010000.00")  # 2 M de compras + 10,000 de inversiones
    # Deudas > créditos: acumulable = 250,000 × factor (INPC aún no publicado: estimación).
    assert ren["aai_acumulable"] == (D("250000") * D(d["aai"]["factor_ajuste"])).quantize(
        D("0.01")) and d["origen_saldos"] == "capturados"
    assert ren["ptu_pagada"] == D("80000.00")
    assert ren["pagos_provisionales"] == sum(
        D(p["pago_efectuado"]) for p in d["pagos_provisionales"])
    assert ren["retenciones"] == D("5000.00")
    assert ren["a_cargo"] - ren["a_favor"] == (
        ren["isr_ejercicio"] - ren["pagos_provisionales"] - ren["retenciones"])
    assert d["coeficiente"]["aplicado"] == "0.1845"
    assert D(d["ptu"]["a_repartir"]) == (D(d["ptu"]["renta_gravable"]) * D("0.10")).quantize(
        D("0.01"))

    x = a.post(f"{IMP}/moral/anual/papeles", json={"anio": 2026})
    assert x.status_code == 200 and x.content[:2] == b"PK"
    libro = load_workbook(io.BytesIO(x.content))
    assert {"Cálculo ISR", "Ingresos", "Deducciones", "Ajuste por inflación", "Coeficiente",
            "PTU", "Pérdidas", "Pagos provisionales", "Notas"} <= set(libro.sheetnames)
    calculo = {f[0]: f[1] for f in libro["Cálculo ISR"].iter_rows(min_row=4, values_only=True)}
    for renglon in d["renglones"]:
        assert D(str(calculo[renglon["concepto"]])) == D(renglon["importe"])
    ingresos = list(libro["Ingresos"].iter_rows(min_row=4, values_only=True))
    total = next(f for f in ingresos if f[0] == "Total")
    assert D(str(total[1])) == ren["ingresos"]
    pagos = list(libro["Pagos provisionales"].iter_rows(min_row=4, values_only=True))
    assert D(str(pagos[2][12])) == D(d["pagos_provisionales"][2]["pago_efectuado"])
    aai = {f[0]: f[1] for f in libro["Ajuste por inflación"].iter_rows(values_only=True)}
    assert D(str(aai["Ajuste anual por inflación acumulable"])) == ren["aai_acumulable"]


def test_anual_pm_desde_la_contabilidad(api, despachos, sesionador):
    db = sesionador()
    c = db.query(Cliente).filter_by(clave="moral").one()
    catalogo.sembrar_base(db, c)
    cs = catalogo.por_numero(db, c)

    def pol(fecha, *movs):
        polizas.crear_manual(db, c, tipo="diario", fecha=fecha, concepto="prueba", uid="u",
                             movimientos=[MovimientoEntrada(cs[n].id, D(x), D(y))
                                          for n, x, y in movs])

    pol(date(2026, 1, 2), ("102.01.01", "100000", "0"), ("301.01", "0", "100000"))
    pol(date(2026, 2, 1), ("115.01", "50000", "0"), ("102.01.01", "0", "50000"))
    pol(date(2026, 3, 10), ("105.01", "116000", "0"), ("401.01", "0", "100000"),
        ("209.01", "0", "16000"))
    pol(date(2026, 3, 11), ("501.01", "40000", "0"), ("115.01", "0", "40000"))
    pol(date(2026, 3, 20), ("601.84", "15000", "0"), ("701.10", "500", "0"),
        ("704.23", "0", "200"), ("102.01.01", "0", "15300"))
    pol(date(2026, 3, 25), ("156.01", "20000", "0"), ("201.01", "0", "20000"))
    pol(date(2026, 4, 5), ("102.01.01", "58000", "0"), ("105.01", "0", "58000"))
    pol(date(2026, 12, 31), ("613.05", "1000", "0"), ("171.03", "0", "1000"))
    db.commit()
    db.close()

    d = api("uid-a").post(f"{IMP}/moral/anual", json={"anio": 2026}).json()
    assert d["fuente"] == "contabilidad" and d["origen_saldos"] == "contabilidad"
    ren = {x["clave"]: D(x["importe"]) for x in d["renglones"]}
    assert ren["ingresos"] == D("100200.00")  # 401.01 + 704.23
    assert ren["deducciones"] == D("55500.00")  # 501.01 + 601.84 + 701.10
    assert sum(D(m["no_deducibles"]) for m in d["meses"]) == D("1000.00")  # 613.05
    # Saldos de fin de mes: créditos 102 + 105; deudas 201 (209 no cuenta, no se ha cobrado).
    a = d["aai"]
    assert a["saldos_creditos"][:4] == ["100000.00", "50000.00", "150700.00", "150700.00"]
    assert a["saldos_deudas"][:4] == ["0.00", "0.00", "20000.00", "20000.00"]
    assert a["promedio_creditos"] == "138083.33" and a["promedio_deudas"] == "16666.67"
    factor = D(a["factor_ajuste"])
    assert D(a["deducible"]) == (D("121416.66") * factor).quantize(D("0.01"))
    clases = {x["numero"]: x["clase"] for x in d["cuentas"]}
    assert clases["613.05"] == "no_deducible" and clases["704.23"] == "ingreso"


def test_resico_pm_por_flujo(api, despachos):
    a = api("uid-a")
    _subir(a, "resico", [_venta(1, 1, "500000.00"), _compra(1, 1, "300000.00")])
    d = a.post(f"{IMP}/resico/pm/pagos-provisionales", json={"anio": 2026, "mes": 1}).json()
    assert d["tipo"] == "resico_pm" and d["coeficiente"] == "0.0000"
    assert d["pagos"][0]["isr_causado"] == "60000.00"  # (500,000 − 300,000) × 30 %
    an = a.post(f"{IMP}/resico/anual", json={"anio": 2026}).json()
    assert an["tipo"] == "resico_pm"
    assert next(x for x in an["renglones"] if x["clave"] == "isr_ejercicio")["importe"] == (
        "60000.00")
    res = a.post(f"{IMP}/resico/resumen",
                 json={"anio": 2026, "mes": 1, "fecha_pago": "2026-02-17"}).json()
    assert next(f for f in res["renglones"]
                if f["clave"] == "isr_resico_pm")["importe_sistema"] == "60000.00"


def test_anual_pf_612_y_papeles_cuadran(api, despachos):
    a = api("uid-a")
    _subir(a, "fisica", [_venta(1, 1, "30000.00", RFC_CLIENTE),
                         _venta(2, 2, "30000.00", RFC_CLIENTE),
                         _compra(1, 1, "5000.00", RFC_CLIENTE)])
    r = a.post(f"{IMP}/fisica/anual", json={"anio": 2026, "pagos_provisionales": {"1": "100"}})
    assert r.status_code == 200, r.text
    d = r.json()
    assert d["tipo"] == "pf" and d["regimen"] == "612" and len(d["meses"]) == 12
    ren = {x["clave"]: D(x["importe"]) for x in d["renglones"]}
    assert ren["ingresos"] == D("60000.00")
    assert d["meses"][0]["pago_efectuado"] == "100.00" and d["meses"][0]["pago_capturado"]
    assert ren["pagos_provisionales"] == sum(D(m["pago_efectuado"]) for m in d["meses"])
    assert ren["a_cargo"] - ren["a_favor"] == (
        ren["isr_ejercicio"] - ren["pagos_provisionales"] - ren["retenciones"])
    x = a.post(f"{IMP}/fisica/anual/papeles", json={"anio": 2026,
                                                    "pagos_provisionales": {"1": "100"}})
    libro = load_workbook(io.BytesIO(x.content))
    assert libro.sheetnames == ["Cálculo ISR", "Mes a mes", "Notas"]
    filas = list(libro["Mes a mes"].iter_rows(min_row=4, values_only=True))
    total = filas[-1]
    assert total[0] == "Total" and D(str(total[1])) == ren["ingresos"]
    assert D(str(total[7])) == ren["pagos_provisionales"]
    calculo = {f[0]: f[1] for f in libro["Cálculo ISR"].iter_rows(min_row=4, values_only=True)}
    for renglon in d["renglones"]:
        assert D(str(calculo[renglon["concepto"]])) == D(renglon["importe"])


def test_errores_de_regimen(api, despachos, sesionador):
    a = api("uid-a")
    assert a.post(f"{IMP}/fisica/pm/pagos-provisionales",
                  json={"anio": 2026, "mes": 1}).status_code == 422
    assert a.get(f"{IMP}/fisica/pm/datos/2026").status_code == 422
    db = sesionador()
    db.query(Cliente).filter_by(clave="fisica").one().datos = {"regimen": "605"}
    db.commit()
    db.close()
    assert a.post(f"{IMP}/fisica/anual", json={"anio": 2026}).status_code == 422
    assert a.put(f"{IMP}/moral/pm/datos/2026",
                 json={"saldos_creditos": ["1"], "saldos_deudas": []}).status_code == 422
    assert a.put(f"{IMP}/moral/pm/datos/2026",
                 json={"retenciones_bancarias": {"13": "1"}}).status_code == 422


def test_multi_tenant(api, despachos):
    a, b = api("uid-a"), api("uid-b")
    a.put(f"{IMP}/moral/pm/datos/2026", json={"coeficiente_utilidad": "0.5"})
    for metodo, ruta, cuerpo in (
        ("get", "/moral/pm/datos/2026", None),
        ("put", "/moral/pm/datos/2026", {}),
        ("post", "/moral/pm/pagos-provisionales", {"anio": 2026, "mes": 1}),
        ("post", "/moral/anual", {"anio": 2026}),
        ("post", "/moral/anual/papeles", {"anio": 2026}),
    ):
        kw = {"json": cuerpo} if cuerpo is not None else {}
        assert getattr(b, metodo)(IMP + ruta, **kw).status_code == 404
    # B tiene su propio 601 con el mismo RFC: sus datos no son los de A.
    assert b.get(f"{IMP}/ajeno/pm/datos/2026").json()["datos"]["coeficiente_utilidad"] is None
    assert a.get(f"{IMP}/moral/pm/datos/2026").json()["datos"]["coeficiente_utilidad"] == "0.5"
    assert a.get(f"{IMP}/ajeno/pm/datos/2026").status_code == 404


def test_calendario_anual_de_persona_moral_en_marzo(api):
    a = api("uid-a")
    def anual(regimen, rfc):
        r = a.post("/api/v1/calendario", json={"contributor_type": "pyme", "regimen": regimen,
                                               "rfc": rfc, "year": 2026})
        assert r.status_code == 200, r.text
        return next(o for o in r.json()["obligaciones"] if o["periodicidad"] == "anual")
    assert anual("601", RFC_PM)["fecha_limite"] == "2027-03-31"
    assert anual("626", RFC_PM)["fecha_limite"] == "2027-03-31"  # RESICO PM (Art. 212)
    assert anual("612", RFC_CLIENTE)["fecha_limite"] == "2027-04-30"
    assert anual("626", RFC_CLIENTE)["fecha_limite"] == "2027-04-30"
