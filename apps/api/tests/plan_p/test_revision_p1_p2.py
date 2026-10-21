"""
Correcciones del revisor a las fases P1 y P2 (contabilidad, Excel, roles, Anexo 24).

Un test (o más) por punto. Los de descarga SAT, base64, portal y semáforo viven junto a sus
módulos (`test_sat_descarga.py`, `test_sat_soap.py`, `test_opinion_cumplimiento.py`).
"""

from __future__ import annotations

import io
from datetime import UTC, date, datetime
from decimal import Decimal

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import delete, select

from app.boveda_cfdi import servicio
from app.boveda_cfdi.calculo import facturas_para_calculo
from app.contabilidad import balanza, catalogo, cierre, config, excel, generador, polizas
from app.contabilidad.anexo24 import generador as anexo24
from app.contabilidad.conciliacion import reporte
from app.contabilidad.conciliacion import servicio as conciliacion
from app.exceptions import FiscalAgentError
from app.exportar.excel_seguro import asegurar_celda, libro_a_bytes, texto_seguro
from app.integraciones.almacenamiento import AlmacenamientoMemoria
from app.modelos.cfdi import Comprobante
from app.modelos.contabilidad import CuentaContable, Poliza
from app.modelos.nucleo import AsignacionCliente, Cliente, Usuario
from app.modelos.sat_verificacion import AlertaRiesgo
from app.tenancy import asegurar_usuario
from tests.plan_p.conta_sintetico import (
    RFC_CLIENTE,
    RFC_PROVEEDOR,
    Concepto,
    factura,
    nomina,
    pago20,
    uuid_n,
)

D = Decimal
OCTUBRE = (date(2026, 10, 1), date(2026, 10, 31))
NOVIEMBRE = (date(2026, 11, 1), date(2026, 11, 30))
MARZO = (date(2026, 3, 1), date(2026, 3, 31))


@pytest.fixture
def cliente(db_sesion):
    asegurar_usuario(db_sesion, "uid-rev")
    c = Cliente(despacho_id="uid-rev", clave="taller", nombre="Taller", rfc=RFC_CLIENTE)
    db_sesion.add(c)
    db_sesion.flush()
    catalogo.sembrar_base(db_sesion, c)
    return c


def subir(db, cliente, *xmls: bytes) -> None:
    rep = servicio.importar(db, cliente, [(f"{i}.xml", io.BytesIO(x)) for i, x in
                                          enumerate(xmls)], almacen=AlmacenamientoMemoria())
    assert not rep.rechazados, rep.rechazados


def poliza_de(db, cliente, clave: str) -> Poliza:
    return db.scalar(select(Poliza).where(Poliza.cliente_id == cliente.id,
                                          Poliza.clave_origen == clave))


def movs(db, p: Poliza) -> dict[str, tuple[Decimal, Decimal]]:
    salida: dict[str, tuple[Decimal, Decimal]] = {}
    for m, c in polizas.movimientos(db, p.id):
        car, abo = salida.get(c.numero, (D(0), D(0)))
        salida[c.numero] = (car + m.cargo, abo + m.abono)
    return salida


def saldo(db, cliente, numero: str, desde: date, hasta: date) -> str:
    b = balanza.balanza(db, cliente, desde, hasta, solo="todas")
    return next(r for r in b["renglones"] if r["numero"] == numero)["saldo_final"]


def con_fecha_pago(xml: bytes, emitido: str, pagado: str) -> bytes:
    """El complemento timbrado `emitido` con su `FechaPago` = `pagado`."""
    return xml.replace(f'FechaPago="{emitido}"'.encode(), f'FechaPago="{pagado}"'.encode())


def relacionada(xml: bytes, uuid: str, tipo: str = "01") -> bytes:
    nodo = (f'<cfdi:CfdiRelacionados TipoRelacion="{tipo}"><cfdi:CfdiRelacionado UUID="{uuid}"/>'
            "</cfdi:CfdiRelacionados>")
    return xml.replace(b"<cfdi:Emisor", nodo.encode() + b"<cfdi:Emisor", 1)


# ── 1. inyección de fórmulas en Excel ──


def test_texto_seguro_y_celdas():
    assert texto_seguro("=1+1") == "'=1+1"
    for peligroso in ("+5", "-2", "@SUM(A1)", "\tx", "\rx"):
        assert texto_seguro(peligroso) == "'" + peligroso
    assert texto_seguro("Venta") == "Venta" and texto_seguro(10) == 10
    wb = Workbook()
    celda = wb.active.cell(1, 1, value='=HYPERLINK("http://malo","clic")')
    assert celda.data_type == "f"
    asegurar_celda(celda)
    assert celda.data_type == "s" and celda.quotePrefix


def _sin_formulas(datos: bytes, texto: str) -> None:
    wb = load_workbook(io.BytesIO(datos))
    encontradas = [c for h in wb.worksheets for fila in h.iter_rows() for c in fila
                   if c.value == texto]
    assert encontradas, "el texto no llegó al Excel"
    for c in encontradas:
        assert c.data_type == "s" and c.quotePrefix  # texto literal, no fórmula
    assert not any(c.data_type == "f" for h in wb.worksheets for f in h.iter_rows() for c in f)


def test_excel_de_polizas_y_conciliacion_sin_formulas():
    malo = '=HYPERLINK("http://malo.example","clic")'
    lista = [{"folio": "D-2026-03-0001", "fecha": "2026-03-01", "estado": "activa",
              "concepto": malo, "total_cargos": "1.00", "total_abonos": "1.00",
              "movimientos": [{"cuenta": "101.01", "nombre_cuenta": "+Caja", "concepto": malo,
                               "uuid": "@x", "cargo": "1.00", "abono": "0.00"}]}]
    _sin_formulas(excel.polizas(lista, {"nombre": "-Cliente", "rfc": "X"}, "a", "b"), malo)
    r = {"cuenta_bancaria": {"nombre": "BBVA", "numero": "1", "cuenta_contable": "102"},
         "corte": "2026-03-31", "saldo_banco": "0", "saldo_libros": "0",
         "saldo_banco_conciliado": "0", "saldo_libros_conciliado": "0", "diferencia": "0",
         "totales": {k: "0" for k, _t in reporte.SECCIONES},
         **{k: [] for k, _t in reporte.SECCIONES}}
    r["abonos_banco_no_registrados"] = [{"fecha": "2026-03-02", "folio": "=1+1",
                                         "concepto": malo, "importe": "5.00"}]
    _sin_formulas(reporte.a_excel(r, "Taller"), malo)


def test_libro_a_bytes_guarda_todo_como_texto():
    wb = Workbook()
    wb.active.append(["=cmd|' /C calc'!A0", 3])
    wb2 = load_workbook(io.BytesIO(libro_a_bytes(wb)))
    assert wb2.active["A1"].data_type == "s" and wb2.active["B1"].value == 3


# ── 2. fecha del complemento de pago = FechaPago ──


def test_pago_del_31_oct_timbrado_el_3_nov_es_de_octubre(db_sesion, cliente):
    fac, pag = uuid_n(1, "F0C00000"), uuid_n(2, "P0C00000")
    subir(db_sesion, cliente,
          factura(fac, metodo="PPD", forma="99", fecha="2026-10-10T10:00:00"),
          con_fecha_pago(pago20(pag, [(fac, "1160.00", "1160.00")],
                                fecha="2026-11-03T09:00:00"),
                         "2026-11-03T09:00:00", "2026-10-31T18:00:00"))
    generador.generar(db_sesion, cliente, *OCTUBRE)
    p = poliza_de(db_sesion, cliente, f"cfdi:{pag}")
    assert p is not None and p.fecha == date(2026, 10, 31)
    assert movs(db_sesion, p)["208.01"] == (0, D("160.00"))
    # Octubre: cobrado y traspasado; noviembre no lo vuelve a generar.
    assert saldo(db_sesion, cliente, "209.01", *OCTUBRE) == "0.00"
    assert saldo(db_sesion, cliente, "208.01", *OCTUBRE) == "-160.00"
    assert saldo(db_sesion, cliente, "105.01", *OCTUBRE) == "0.00"
    r = generador.generar(db_sesion, cliente, *NOVIEMBRE)
    assert r.generadas == 0
    # Cálculos (facturas-calculo): el pago cae en octubre con fecha 31-oct, no en noviembre.
    oct_ = facturas_para_calculo(db_sesion, cliente, *OCTUBRE)
    pagos = [f for f in oct_.facturas if f.tipo == "P"]
    assert [(f.fecha, f.monto_pago) for f in pagos] == [(date(2026, 10, 31), 1160.0)]
    assert not [f for f in facturas_para_calculo(db_sesion, cliente, *NOVIEMBRE).facturas
                if f.tipo == "P"]
    # Conciliación de octubre: el depósito del 31 ya no queda en tránsito.
    cb = conciliacion.crear_cuenta(db_sesion, cliente, {
        "banco": "012", "nombre": "BBVA", "numero": "0123456789", "clabe": "",
        "cuenta_contable_id": catalogo.por_numero(db_sesion, cliente)["102.01.01"].id})
    csv = ("Fecha,Descripción,Referencia,Retiros,Depósitos,Saldo\n"
           'Saldo inicial,,,,,"0.00"\n'
           '31/10/2026,SPEI RECIBIDO,REF1,,"1,160.00","1,160.00"\n')
    conciliacion.importar(db_sesion, cliente, cb, nombre="oct.csv", datos=csv.encode())
    assert conciliacion.conciliar_automatico(db_sesion, cliente, cb, None, None)[
        "conciliados"] == 1
    rep = reporte.reporte(db_sesion, cliente, cb, date(2026, 10, 31))
    assert rep["depositos_en_transito"] == [] and rep["abonos_banco_no_registrados"] == []
    assert rep["cuadra"] is True


def test_complemento_con_pagos_de_fechas_distintas_se_parte(db_sesion, cliente):
    fac, pag = uuid_n(3, "F0C00000"), uuid_n(4, "P0C00000")
    xml = pago20(pag, [(fac, "1160.00", "580.00")], fecha="2026-11-03T09:00:00")
    # Dos pagos en un complemento: 31-oct y 2-nov, de 580 cada uno.
    texto = xml.decode()
    i, j = texto.index("<pago20:Pago "), texto.index("</pago20:Pagos>")
    nodo = texto[i:j]
    segundo = nodo.replace('ImpSaldoAnt="1160.00"', 'ImpSaldoAnt="580.00"').replace(
        'ImpSaldoInsoluto="580.00"', 'ImpSaldoInsoluto="0.00"').replace(
        'NumParcialidad="1"', 'NumParcialidad="2"')
    primero = nodo.replace("2026-11-03T09:00:00", "2026-10-31T10:00:00")
    segundo = segundo.replace("2026-11-03T09:00:00", "2026-11-02T10:00:00")
    xml = (texto[:i] + primero + segundo + texto[j:]).encode()
    subir(db_sesion, cliente, factura(fac, metodo="PPD", forma="99",
                                      fecha="2026-10-10T10:00:00"), xml)
    generador.generar(db_sesion, cliente, *OCTUBRE)
    p_oct = poliza_de(db_sesion, cliente, f"cfdi:{pag}@2026-10-31")
    assert p_oct.fecha == date(2026, 10, 31)
    assert movs(db_sesion, p_oct)["102.01.01"] == (D("580.00"), 0)
    assert poliza_de(db_sesion, cliente, f"cfdi:{pag}@2026-11-02") is None
    generador.generar(db_sesion, cliente, *NOVIEMBRE)
    p_nov = poliza_de(db_sesion, cliente, f"cfdi:{pag}@2026-11-02")
    assert p_nov.fecha == date(2026, 11, 2)
    assert movs(db_sesion, p_nov)["208.01"] == (0, D("80.00"))
    assert saldo(db_sesion, cliente, "209.01", *NOVIEMBRE) == "0.00"
    # Cancelar el complemento cancela sus dos pólizas.
    comp = db_sesion.scalar(select(Comprobante).where(Comprobante.uuid == pag))
    assert polizas.cancelar_por_cfdi(db_sesion, comp) == "cancelada"
    assert {p_oct.estado, p_nov.estado} == {"cancelada"}
    assert generador.generar(db_sesion, cliente, *OCTUBRE).generadas == 0


# ── 4. nota de crédito sobre PPD con saldo ──


def test_nota_de_credito_sobre_ppd_revierte_contra_no_cobrado(db_sesion, cliente):
    fac, nc, pag = uuid_n(1, "F4000000"), uuid_n(2, "E4000000"), uuid_n(3, "P4000000")
    subir(db_sesion, cliente, factura(fac, metodo="PPD", forma="99"),
          relacionada(factura(nc, tipo="E", metodo="PUE", forma="03",
                              conceptos=[Concepto("100.00")]), fac))
    generador.generar(db_sesion, cliente, *MARZO)
    m = movs(db_sesion, poliza_de(db_sesion, cliente, f"cfdi:{nc}"))
    assert m["209.01"] == (D("16.00"), 0) and "208.01" not in m  # no contra cobrado
    assert saldo(db_sesion, cliente, "209.01", *MARZO) == "-144.00"
    # El pago del saldo (1044) liquida: traspasa exactamente lo que queda (144).
    subir(db_sesion, cliente, pago20(pag, [(fac, "1044.00", "1044.00")],
                                     fecha="2026-03-25T10:00:00"))
    generador.generar(db_sesion, cliente, *MARZO)
    assert movs(db_sesion, poliza_de(db_sesion, cliente, f"cfdi:{pag}"))["208.01"] == (
        0, D("144.00"))
    assert saldo(db_sesion, cliente, "209.01", *MARZO) == "0.00"


def test_nota_recibida_sobre_ppd_revierte_contra_iva_pendiente(db_sesion, cliente):
    fac, nc = uuid_n(4, "F4000000"), uuid_n(5, "E4000000")
    subir(db_sesion, cliente,
          factura(fac, emisor=RFC_PROVEEDOR, receptor=RFC_CLIENTE, metodo="PPD", forma="99"),
          relacionada(factura(nc, tipo="E", emisor=RFC_PROVEEDOR, receptor=RFC_CLIENTE,
                              metodo="PUE", forma="03", conceptos=[Concepto("100.00")]),
                      fac, "03"))
    generador.generar(db_sesion, cliente, *MARZO)
    m = movs(db_sesion, poliza_de(db_sesion, cliente, f"cfdi:{nc}"))
    assert m["119.01"] == (0, D("16.00")) and "118.01" not in m
    assert saldo(db_sesion, cliente, "119.01", *MARZO) == "144.00"


# ── 5. residuos de parcialidades ──


def test_tres_parcialidades_dejan_cero_en_209(db_sesion, cliente):
    fac = uuid_n(1, "F5000000")
    p1, p2, p3 = (uuid_n(n, "P5000000") for n in (1, 2, 3))
    subir(db_sesion, cliente, factura(fac, metodo="PPD", forma="99"),
          pago20(p1, [(fac, "1160.00", "386.67")], fecha="2026-03-11T10:00:00"),
          pago20(p2, [(fac, "773.33", "386.67")], fecha="2026-03-12T10:00:00"),
          pago20(p3, [(fac, "386.66", "386.66")], fecha="2026-03-13T10:00:00"))
    generador.generar(db_sesion, cliente, *MARZO)
    ivas = [movs(db_sesion, poliza_de(db_sesion, cliente, f"cfdi:{u}"))["208.01"][1]
            for u in (p1, p2, p3)]
    # Proporcional: 53.33 + 53.33 + 53.33 = 159.99. La última (saldo insoluto 0) lleva el
    # resto: 53.34, y 209 queda en cero.
    assert ivas == [D("53.33"), D("53.33"), D("53.34")]
    assert saldo(db_sesion, cliente, "209.01", *MARZO) == "0.00"
    assert saldo(db_sesion, cliente, "208.01", *MARZO) == "-160.00"


# ── 6. retenciones en PPD ──


def test_retenciones_ppd_quedan_pendientes_hasta_el_pago(db_sesion, cliente):
    fac, pag = uuid_n(1, "F6000000"), uuid_n(2, "P6000000")
    neto = D("10000.00") + D("1600.00") - D("1000.00") - D("1066.67")
    subir(db_sesion, cliente,
          factura(fac, metodo="PPD", forma="99",
                  conceptos=[Concepto("10000.00", ret_isr="1000.00", ret_iva="1066.67")]),
          pago20(pag, [(fac, f"{neto}", f"{neto}")], fecha="2026-03-20T10:00:00"))
    generador.generar(db_sesion, cliente, *MARZO)
    m = movs(db_sesion, poliza_de(db_sesion, cliente, f"cfdi:{fac}"))
    assert m["113.08.01"] == (D("1066.67"), 0) and m["113.08.02"] == (D("1000.00"), 0)
    assert "113.01" not in m and "113.02" not in m
    mp = movs(db_sesion, poliza_de(db_sesion, cliente, f"cfdi:{pag}"))
    assert mp["113.01"] == (D("1066.67"), 0) and mp["113.08.01"] == (0, D("1066.67"))
    assert mp["113.02"] == (D("1000.00"), 0) and mp["113.08.02"] == (0, D("1000.00"))
    assert saldo(db_sesion, cliente, "113.08.01", *MARZO) == "0.00"


def test_retenciones_recibidas_y_parciales_proporcionales(db_sesion, cliente):
    fac, pag = uuid_n(3, "F6000000"), uuid_n(4, "P6000000")
    total = D("5000.00") + D("800.00") - D("500.00") - D("533.33")  # 4766.67
    subir(db_sesion, cliente,
          factura(fac, emisor=RFC_PROVEEDOR, receptor=RFC_CLIENTE, metodo="PPD", forma="99",
                  conceptos=[Concepto("5000.00", clave="80131500", ret_isr="500.00",
                                      ret_iva="533.33")]),
          pago20(pag, [(fac, f"{total}", "2383.34")], emisor=RFC_PROVEEDOR,
                 receptor=RFC_CLIENTE, fecha="2026-03-21T10:00:00"))
    generador.generar(db_sesion, cliente, *MARZO)
    mp = movs(db_sesion, poliza_de(db_sesion, cliente, f"cfdi:{pag}"))
    # 2383.34 / 4766.67 = 50 %: la mitad de cada retención pasa a 216.03 / 216.10.
    assert mp["216.12.03"] == (D("250.00"), 0) and mp["216.03"] == (0, D("250.00"))
    assert mp["216.12.01"] == (D("266.67"), 0) and mp["216.10"] == (0, D("266.67"))
    assert saldo(db_sesion, cliente, "216.12.03", *MARZO) == "-250.00"


def test_retenciones_al_pago_apagado_es_el_comportamiento_anterior(db_sesion, cliente):
    config.guardar(db_sesion, cliente, {"retenciones_al_pago": False})
    assert config.como_dict(db_sesion, cliente)["retenciones_al_pago"] is False
    fac = uuid_n(5, "F6000000")
    subir(db_sesion, cliente, factura(
        fac, metodo="PPD", forma="99",
        conceptos=[Concepto("10000.00", ret_isr="1000.00", ret_iva="1066.67")]))
    generador.generar(db_sesion, cliente, *MARZO)
    m = movs(db_sesion, poliza_de(db_sesion, cliente, f"cfdi:{fac}"))
    assert m["113.01"] == (D("1066.67"), 0) and "113.08.01" not in m


def test_catalogo_viejo_recibe_las_cuentas_pendientes_al_usarlas(db_sesion, cliente):
    viejas = ("216.12.01", "216.12.02", "216.12.03", "216.12")
    for numero in viejas:  # cliente sembrado antes de que existieran
        db_sesion.execute(delete(CuentaContable).where(CuentaContable.cliente_id == cliente.id,
                                                       CuentaContable.numero == numero))
    db_sesion.flush()
    fac = uuid_n(6, "F6000000")
    subir(db_sesion, cliente, factura(
        fac, emisor=RFC_PROVEEDOR, receptor=RFC_CLIENTE, metodo="PPD", forma="99",
        conceptos=[Concepto("5000.00", ret_isr="500.00", ret_iva="533.33")]))
    r = generador.generar(db_sesion, cliente, *MARZO)
    assert r.errores == [] and r.generadas == 1
    cuentas = catalogo.por_numero(db_sesion, cliente)
    assert not cuentas["216.12"].afectable and cuentas["216.12.02"].afectable
    assert movs(db_sesion, poliza_de(db_sesion, cliente, f"cfdi:{fac}"))["216.12.02"] == (
        0, D("500.00"))


# ── 7. nómina contable ──


def test_nomina_separa_imss_otras_deducciones_y_otros_pagos(db_sesion, cliente):
    n = uuid_n(7, "N7000000")
    subir(db_sesion, cliente, nomina(n, otros_pagos="50.00",
                                     mas_deducciones=(("004", "100.00"),),
                                     mas_otros_pagos=(("003", "300.00"), ("001", "20.00"))))
    generador.generar(db_sesion, cliente, *MARZO)
    p = poliza_de(db_sesion, cliente, f"cfdi:{n}")
    assert p.diferencia == 0
    neto = D("10000.00") + D("370.00") - D("1350.00")
    assert movs(db_sesion, p) == {
        "601.01": (D("10000.00"), 0),
        "110.01": (D("50.00"), 0),  # subsidio (002)
        "107.01": (D("300.00"), 0),  # viáticos (003)
        "216.01": (D("20.00"), D("1000.00")),  # reintegro ISR (001) / ISR retenido
        "216.11": (0, D("250.00")),  # cuota obrera IMSS, agrupador 216.11
        "205.06": (0, D("100.00")),  # otra deducción (préstamo 004)
        "102.01.01": (0, neto),
    }


# ── 10. reversa de cancelación fallida → alerta ──


def test_reversa_imposible_deja_alerta_consultable(db_sesion, cliente):
    from app.sat_verificacion import estatus

    a = uuid_n(1, "A1000000")
    subir(db_sesion, cliente, factura(a))
    generador.generar(db_sesion, cliente, *MARZO)
    hoy = date.today()
    polizas.cerrar_periodo(db_sesion, cliente, 2026, 3, "u")
    polizas.cerrar_periodo(db_sesion, cliente, hoy.year, hoy.month, "u")  # la reversa tampoco
    comp = db_sesion.scalar(select(Comprobante).where(Comprobante.uuid == a))

    class Est:
        encontrado, resumen, en_proceso_de_cancelacion = True, "cancelado", False
        codigo_estatus = estado = es_cancelable = estatus_cancelacion = "x"
        validacion_efos = ""

    r = estatus.ResultadoVerificacion()
    estatus._aplicar(db_sesion, comp, Est(), datetime.now(UTC), r)  # type: ignore[arg-type]
    assert comp.estatus_sat == "cancelado" and r.reversas_fallidas == 1
    alerta = db_sesion.scalar(select(AlertaRiesgo).where(AlertaRiesgo.cliente_id == cliente.id))
    assert alerta.tipo == polizas.TIPO_ALERTA_REVERSA and alerta.detalle["uuid"] == a
    assert "cerrado" in alerta.detalle["error"]
    # Idempotente: una sola alerta por CFDI.
    assert polizas.cancelar_por_cfdi_o_alertar(db_sesion, comp) == "alerta"
    assert len(db_sesion.scalars(select(AlertaRiesgo)).all()) == 1


# ── 11. roles ──


@pytest.fixture
def despacho_con_auxiliar(sesionador):
    db = sesionador()
    asegurar_usuario(db, "uid-a")
    c = Cliente(despacho_id="uid-a", clave="taller", nombre="Taller", rfc=RFC_CLIENTE)
    db.add(c)
    db.flush()
    catalogo.sembrar_base(db, c)
    db.add(Usuario(uid="aux", despacho_id="uid-a", rol="auxiliar"))
    db.flush()
    db.add(AsignacionCliente(usuario_uid="aux", cliente_id=c.id))
    db.commit()
    cuenta = catalogo.por_numero(db, c)["102.01.01"].id
    db.close()
    return cuenta


def test_auxiliar_solo_lee_genera_e_importa(api, despacho_con_auxiliar):
    aux, duena = api("aux"), api("uid-a")
    conta = "/api/v1/contabilidad/clientes/taller"
    banco = "/api/v1/conciliacion/clientes/taller"
    ce = "/api/v1/contabilidad-electronica/clientes/taller"
    # Lectura y lo permitido.
    assert aux.get(f"{conta}/cuentas").status_code == 200
    assert aux.get(f"{conta}/config").status_code == 200
    assert aux.post(f"{conta}/polizas/generar",
                    json={"desde": "2026-03-01", "hasta": "2026-03-31"}).status_code == 200
    cb = duena.post(f"{banco}/cuentas", json={
        "banco": "012", "nombre": "BBVA", "numero": "1", "clabe": "",
        "cuenta_contable_id": despacho_con_auxiliar})
    assert cb.status_code == 201, cb.text
    cb_id = cb.json()["cuenta"]["id"]
    csv = b'Fecha,Descripcion,Referencia,Retiros,Depositos,Saldo\n02/03/2026,X,R,,"10.00","10.00"\n'
    imp = aux.post(f"{banco}/cuentas/{cb_id}/estados", files={"archivo": ("m.csv", csv)})
    assert imp.status_code == 200, imp.text
    # Escritura: 403.
    prohibidas = [
        ("put", f"{conta}/config", {"pue_registra_cobro": False}),
        ("post", f"{conta}/cuentas/sembrar", None),
        ("post", f"{conta}/reglas", {"criterio": "rfc", "valor": "X", "cuenta_numero": "601.84"}),
        ("post", f"{conta}/polizas", {"tipo": "diario", "fecha": "2026-03-02", "concepto": "x",
                                      "movimientos": [
                                          {"cuenta_id": despacho_con_auxiliar, "cargo": "1"},
                                          {"cuenta_id": despacho_con_auxiliar, "abono": "1"}]}),
        ("post", f"{conta}/periodos/2026/3/cerrar", {"forzar": False, "motivo": ""}),
        ("post", f"{conta}/cierre-anual/2026", {}),
        ("post", f"{banco}/cuentas", {"banco": "012", "nombre": "B", "numero": "2", "clabe": "",
                                      "cuenta_contable_id": despacho_con_auxiliar}),
        ("post", f"{banco}/cuentas/{cb_id}/conciliar", {}),
        ("delete", f"{banco}/cuentas/{cb_id}/conciliaciones/1", None),
        ("post", f"{ce}/validar", {"documento": "catalogo", "anio": 2026, "mes": 3}),
        ("post", f"{ce}/zip", {"documento": "catalogo", "anio": 2026, "mes": 3}),
    ]
    for metodo, url, cuerpo in prohibidas:
        kw = {"json": cuerpo} if cuerpo is not None else {}
        r = getattr(aux, metodo)(url, **kw)
        assert r.status_code == 403, (metodo, url, r.status_code, r.text)
    # La dueña sí puede.
    assert duena.put(f"{conta}/config", json={"pue_registra_cobro": False}).status_code == 200


# ── 14. Anexo 24: balanza de cierre (mes 13) ──


def test_balanza_mes_13_con_las_polizas_de_cierre(db_sesion, cliente):
    from lxml import etree

    cliente.rfc = "XAXX010101AB1"
    subir(db_sesion, cliente, factura(uuid_n(1, "B1300000"), emisor=cliente.rfc,
                                      fecha="2026-03-10T10:00:00"))
    generador.generar(db_sesion, cliente, *MARZO)
    sin_cierre = anexo24.Solicitud("balanza", 2026, 13)
    with pytest.raises(FiscalAgentError) as e:
        anexo24.generar(db_sesion, cliente, sin_cierre)
    assert e.value.status_code == 409
    cierre.cierre_anual(db_sesion, cliente, 2026, "u", forzar=True, motivo="prueba")
    doc = anexo24.generar(db_sesion, cliente, sin_cierre)
    assert doc.valido, doc.errores  # valida contra el XSD oficial (Mes 13 está permitido)
    assert doc.nombre.endswith("202613BN")
    raiz = etree.fromstring(doc.xml)
    assert raiz.get("Mes") == "13"
    ctas = {c.get("NumCta"): c for c in raiz}
    # Ingresos: saldo al 31-dic antes del cierre, en cero después.
    assert ctas["401.01"].get("SaldoIni") == "1000.00"
    assert ctas["401.01"].get("Debe") == "1000.00" and ctas["401.01"].get("SaldoFin") == "0.00"
    # La utilidad llega a capital; el balance no se salda (no entra la póliza de balance).
    assert ctas["305.01"].get("SaldoFin") == "1000.00"
    assert ctas["102.01.01"].get("SaldoFin") == "1160.00"
    with pytest.raises(FiscalAgentError, match="sólo existe para la balanza"):
        anexo24.generar(db_sesion, cliente, anexo24.Solicitud("catalogo", 2026, 13))
