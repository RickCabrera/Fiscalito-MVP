"""
P5-03: cuentas por cobrar del despacho — igualas mensuales idempotentes, servicios, pagos
(con CFDI de la bóveda opcional), saldo y antigüedad FIFO, estado de cuenta PDF/Excel,
recordatorios con adaptadores falsos y multi-tenant. "Listo cuando": el despacho ve qué le
debe cada cliente y desde cuándo.
"""

from __future__ import annotations

import io
from datetime import UTC, date, datetime, timedelta
from decimal import Decimal

import pytest
from openpyxl import load_workbook
from sqlalchemy import func, select

from app import almacen_cartera
from app.cobranza import recordatorios
from app.cobranza.servicio import calcular_saldo, generar_igualas
from app.integraciones import registro
from app.modelos.cfdi import Comprobante
from app.modelos.cobranza import CargoCobranza, PagoCobranza
from app.modelos.nucleo import Cliente, Despacho, Usuario
from app.modelos.trabajos import Trabajo
from app.tenancy import asegurar_usuario
from app.trabajos.cola import procesar_siguiente

A, B = "uid-a", "uid-b"
HOY = date(2026, 10, 19)
URL = "/api/v1/cobranza"


@pytest.fixture
def adaptadores():
    registro.limpiar_cache()
    return registro.obtener("correo"), registro.obtener("whatsapp")


def _despacho(sesionador, uid: str = A, clientes=("taller", "tienda")):
    with sesionador() as db:
        asegurar_usuario(db, uid)
        db.get(Despacho, uid).plan = "despacho"
        db.get(Despacho, uid).nombre = "Despacho Pérez"
        for clave in clientes:
            almacen_cartera.crear_cliente(db, uid, clave, {"nombre": clave.title(),
                                                          "rfc": "TAL010101AB1"}, con_tope=False)
        db.commit()


def _honorarios(api, clave="taller", uid=A, **extra):
    cuerpo = {"iguala_mensual": "5000", "dia_cargo": 5, "dias_credito": 10,
              "correo": "pagos@taller.mx", "telefono": "+5215555555555", **extra}
    r = api(uid).put(f"{URL}/clientes/{clave}/honorarios", json=cuerpo)
    assert r.status_code == 200, r.text


def _cargo(fecha: date, importe: str, i: int, dias: int = 10) -> CargoCobranza:
    return CargoCobranza(id=i, tipo="servicio", concepto=f"c{i}", fecha=fecha,
                         vencimiento=fecha + timedelta(days=dias), importe=Decimal(importe))


def _pago(importe: str) -> PagoCobranza:
    return PagoCobranza(fecha=HOY, importe=Decimal(importe))


def test_antiguedad_fifo_por_tramos():
    cargos = [_cargo(HOY - timedelta(days=100), "1000", 1),
              _cargo(HOY - timedelta(days=70), "1000", 2),
              _cargo(HOY - timedelta(days=40), "1000", 3),
              _cargo(HOY - timedelta(days=5), "1000", 4)]
    s = calcular_saldo(cargos, [_pago("1500")], HOY)
    # El pago liquida el de hace 100 días y la mitad del de 70.
    assert s.saldo == Decimal("2500.00")
    assert s.tramos == {"0_30": Decimal("1000"), "31_60": Decimal("1000"),
                        "61_90": Decimal("500"), "mas_90": Decimal("0")}
    assert s.desde == HOY - timedelta(days=70)
    # Vencido: los tres viejos (vencen a los 10 días); el de hace 5 aún no.
    assert s.vencido == Decimal("1500") and s.vencido_desde == HOY - timedelta(days=60)
    # Bordes de tramo: 30 → 0-30, 31 → 31-60, 90 → 61-90, 91 → +90.
    for dias, tramo in ((30, "0_30"), (31, "31_60"), (60, "31_60"), (61, "61_90"),
                        (90, "61_90"), (91, "mas_90")):
        s = calcular_saldo([_cargo(HOY - timedelta(days=dias), "1", 1)], [], HOY)
        assert s.tramos[tramo] == Decimal("1"), dias
    # Pago de más: saldo a favor, sin tramos.
    s = calcular_saldo([_cargo(HOY, "100", 1)], [_pago("150")], HOY)
    assert s.saldo == Decimal("-50.00") and sum(s.tramos.values()) == 0 and s.desde is None


def test_cargos_mensuales_idempotentes(api, sesionador):
    _despacho(sesionador)
    _honorarios(api)
    _honorarios(api, "tienda", dia_cargo=25)
    with sesionador() as db:
        # Día 10: sólo el taller (día 5) ya tocó; la tienda (día 25) todavía no.
        assert len(generar_igualas(db, date(2026, 10, 10))) == 1
        assert generar_igualas(db, date(2026, 10, 10)) == []
        assert len(generar_igualas(db, date(2026, 10, 26))) == 1
        assert generar_igualas(db, date(2026, 10, 31)) == []
        db.commit()
        cargos = list(db.scalars(select(CargoCobranza).order_by(CargoCobranza.id)))
    assert [(c.clave, c.fecha, c.vencimiento, c.importe) for c in cargos] == [
        ("iguala:2026-10", date(2026, 10, 5), date(2026, 10, 15), Decimal("5000.00")),
        ("iguala:2026-10", date(2026, 10, 25), date(2026, 11, 4), Decimal("5000.00"))]


def test_el_trabajo_programado_genera_una_vez(api, sesionador):
    from app.cobranza.trabajos import encolar_cargos
    from app.trabajos.manejadores import cargar_manejadores

    cargar_manejadores()
    _despacho(sesionador)
    _honorarios(api)
    ranura = datetime(2026, 10, 19, 12, 0, tzinfo=UTC)
    for _ in range(2):  # el trabajo corre dos veces (reintento): mismo resultado
        with sesionador() as db:
            encolar_cargos(db, ranura)
            db.commit()
        procesar_siguiente(sesionador, ahora=ranura)
    with sesionador() as db:
        assert db.scalar(select(func.count()).select_from(CargoCobranza)) == 1
        assert {t.estado for t in db.scalars(select(Trabajo))} == {"hecho"}


def test_generar_manual_saldo_y_desde_cuando(api, sesionador):
    _despacho(sesionador)
    _honorarios(api)
    a = api(A)
    for periodo in ("2026-07", "2026-08", "2026-09", "2026-10"):
        assert a.post(f"{URL}/cargos/generar", json={"periodo": periodo}).json()["cargos"] == 1
    assert a.post(f"{URL}/cargos/generar", json={"periodo": "2026-10"}).json()["cargos"] == 0
    r = a.post(f"{URL}/clientes/taller/cargos", json={
        "concepto": "Declaración anual", "importe": "1200", "fecha": "2026-10-01",
        "facturar": True})
    assert r.status_code == 200 and r.json()["cargo"]["vencimiento"] == "2026-10-11"
    assert a.post(f"{URL}/clientes/taller/pagos", json={
        "fecha": "2026-08-20", "importe": "7000", "forma_pago": "03"}).status_code == 200
    d = a.get(f"{URL}/clientes/taller", params={"hoy": HOY.isoformat()}).json()
    # 4 igualas (20,000) + 1,200 - 7,000 = 14,200. El pago cubre julio y 2,000 de agosto.
    assert d["saldo"] == "14200.00"
    assert d["desde"] == "2026-08-05"
    # Al 19/oct: agosto 5 (75 días) → 61-90; sep 5 (44) → 31-60; oct 1 y oct 5 → 0-30.
    assert d["antiguedad"] == {"0_30": "6200.00", "31_60": "5000.00", "61_90": "3000.00",
                               "mas_90": "0.00"}
    assert d["vencido"] == "14200.00" and d["vencido_desde"] == "2026-08-15"
    assert d["movimientos"][-1]["saldo"] == "14200.00"
    res = a.get(f"{URL}/resumen", params={"hoy": HOY.isoformat()}).json()
    assert res["clientes"][0]["cliente"] == "taller" and res["totales"]["saldo"] == "14200.00"
    assert res["clientes"][1]["saldo"] == "0.00"  # la tienda no debe nada


def test_gancho_de_facturacion(api, sesionador):
    _despacho(sesionador)
    a = api(A)
    cargo = a.post(f"{URL}/clientes/taller/cargos", json={
        "concepto": "Extra", "importe": "100", "fecha": "2026-10-01"}).json()["cargo"]
    r = a.put(f"{URL}/clientes/taller/cargos/{cargo['id']}/factura",
              json={"facturar": True, "uuid_cfdi": "aaaaaaaa-bbbb-cccc-dddd-eeeeeeeeeeee"})
    assert r.json()["cargo"]["uuid_cfdi"] == "AAAAAAAA-BBBB-CCCC-DDDD-EEEEEEEEEEEE"
    # Facturado no se cancela sin cancelar antes el CFDI.
    assert a.delete(f"{URL}/clientes/taller/cargos/{cargo['id']}").status_code == 409


def test_pago_ligado_a_cfdi_de_la_boveda(api, sesionador):
    _despacho(sesionador, clientes=("taller", "propio"))
    with sesionador() as db:
        propio = db.scalar(select(Cliente).where(Cliente.clave == "propio"))
        db.add(Comprobante(cliente_id=propio.id, uuid="11111111-2222-3333-4444-555555555555",
                           version="4.0", tipo="I", fecha=datetime(2026, 10, 1),
                           rfc_emisor="DES010101AB1", rfc_receptor="TAL010101AB1",
                           direccion="emitido", total=Decimal("5800"),
                           ruta_xml="x/1.xml", sha256="0" * 64))
        db.commit()
    a = api(A)
    ok = a.post(f"{URL}/clientes/taller/pagos", json={
        "fecha": "2026-10-10", "importe": "5800",
        "cfdi_uuid": "11111111-2222-3333-4444-555555555555"})
    assert ok.status_code == 200, ok.text
    malo = a.post(f"{URL}/clientes/taller/pagos", json={
        "fecha": "2026-10-10", "importe": "1", "cfdi_uuid": "99999999-2222-3333-4444-555555555555"})
    assert malo.status_code == 422
    # Sin CFDI también se registra (es opcional).
    assert a.post(f"{URL}/clientes/taller/pagos", json={
        "fecha": "2026-10-10", "importe": "1"}).status_code == 200


def test_estado_de_cuenta_pdf_y_excel(api, sesionador):
    _despacho(sesionador)
    _honorarios(api)
    a = api(A)
    a.post(f"{URL}/cargos/generar", json={"periodo": "2026-09"})
    a.post(f"{URL}/clientes/taller/pagos", json={"fecha": "2026-09-20", "importe": "2000"})
    pdf = a.get(f"{URL}/clientes/taller/estado-cuenta", params={"formato": "pdf"})
    assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF")
    xlsx = a.get(f"{URL}/clientes/taller/estado-cuenta", params={"formato": "xlsx"})
    assert xlsx.status_code == 200
    hoja = load_workbook(io.BytesIO(xlsx.content))["Estado de cuenta"]
    valores = [c for fila in hoja.iter_rows(values_only=True) for c in fila]
    assert "Saldo" in valores and Decimal("3000") in [Decimal(str(v)) for v in valores
                                                       if isinstance(v, int | float)]


def test_recordatorio_manual_con_adaptador_falso(api, sesionador, adaptadores):
    correo, whatsapp = adaptadores
    _despacho(sesionador)
    _honorarios(api)
    a = api(A)
    a.post(f"{URL}/cargos/generar", json={"periodo": "2026-09"})
    r = a.post(f"{URL}/clientes/taller/recordatorios", json={"canal": "correo"})
    assert r.status_code == 200 and r.json()["recordatorio"]["estado"] == "enviado"
    assert correo.enviados[-1].destinatarios == ["pagos@taller.mx"]
    assert "$5,000.00" in correo.enviados[-1].html
    r = a.post(f"{URL}/clientes/taller/recordatorios", json={"canal": "whatsapp"})
    assert r.json()["recordatorio"]["estado"] == "enviado"
    assert whatsapp.enviados[-1]["telefono"] == "+5215555555555"
    # Sin teléfono capturado: error claro, nada se manda.
    _honorarios(api, "tienda", telefono="")
    assert a.post(f"{URL}/clientes/tienda/recordatorios",
                  json={"canal": "whatsapp"}).status_code == 422


def test_recordatorio_programado_corre_en_la_cola(api, sesionador, adaptadores):
    from app.trabajos.manejadores import cargar_manejadores

    cargar_manejadores()
    correo, _ = adaptadores
    _despacho(sesionador)
    _honorarios(api)
    a = api(A)
    a.post(f"{URL}/cargos/generar", json={"periodo": "2026-09"})
    cuando = datetime.now(UTC) + timedelta(days=2)
    r = a.post(f"{URL}/clientes/taller/recordatorios",
               json={"canal": "correo", "programado_para": cuando.isoformat()})
    assert r.json()["recordatorio"]["estado"] == "programado" and correo.enviados == []
    assert procesar_siguiente(sesionador, ahora=datetime.now(UTC)) is None  # aún no toca
    t = procesar_siguiente(sesionador, ahora=cuando + timedelta(minutes=1))
    assert t.estado == "hecho" and len(correo.enviados) == 1
    detalle = a.get(f"{URL}/clientes/taller").json()
    assert detalle["recordatorios"][0]["estado"] == "enviado"


def test_recordatorios_automaticos(api, sesionador, adaptadores):
    correo, _ = adaptadores
    _despacho(sesionador)
    _honorarios(api, recordatorio_automatico=True, dias_entre_recordatorios=7)
    _honorarios(api, "tienda", recordatorio_automatico=True)
    a = api(A)
    a.post(f"{URL}/cargos/generar", json={"periodo": "2026-09"})  # vence 15/sep
    with sesionador() as db:  # la tienda pagó: no se le molesta
        tienda = db.scalar(select(Cliente).where(Cliente.clave == "tienda"))
        db.add(PagoCobranza(despacho_id=A, cliente_id=tienda.id, fecha=HOY,
                            importe=Decimal("5000")))
        db.commit()
    hoy = date.today()
    with sesionador() as db:
        assert recordatorios.automaticos(db, hoy) == {"enviados": 1, "omitidos": 0}
        db.commit()
    with sesionador() as db:  # el mismo día otra vez: ya se mandó
        assert recordatorios.automaticos(db, hoy) == {"enviados": 0, "omitidos": 1}
    assert len(correo.enviados) == 1


def test_solo_dueno_y_contador(api, sesionador):
    _despacho(sesionador)
    with sesionador() as db:
        db.add(Usuario(uid="uid-aux", despacho_id=A, rol="auxiliar"))
        db.add(Usuario(uid="uid-cont", despacho_id=A, rol="contador"))
        db.commit()
    assert api("uid-aux").get(f"{URL}/resumen").status_code == 403
    assert api("uid-aux").put(f"{URL}/clientes/taller/honorarios",
                              json={"iguala_mensual": "1"}).status_code == 403
    assert api("uid-cont").get(f"{URL}/resumen").status_code == 200


def test_multi_tenant(api, sesionador):
    _despacho(sesionador)
    _despacho(sesionador, uid=B, clientes=("ajeno",))
    _honorarios(api, "ajeno", uid=B)
    b = api(B)
    b.post(f"{URL}/cargos/generar", json={"periodo": "2026-09"})
    detalle_b = b.get(f"{URL}/clientes/ajeno")
    assert detalle_b.status_code == 200, detalle_b.text
    cargo_id = detalle_b.json()["cargos"][0]["id"]
    a = api(A)
    assert a.get(f"{URL}/clientes/ajeno").status_code == 404
    assert a.post(f"{URL}/clientes/ajeno/pagos",
                  json={"fecha": "2026-10-01", "importe": "1"}).status_code == 404
    assert a.delete(f"{URL}/clientes/taller/cargos/{cargo_id}").status_code == 404
    assert [c["cliente"] for c in a.get(f"{URL}/resumen").json()["clientes"]] == [
        "taller", "tienda"]
    # Generar en A no toca las igualas de B.
    assert a.post(f"{URL}/cargos/generar", json={"periodo": "2026-10"}).json()["cargos"] == 0
