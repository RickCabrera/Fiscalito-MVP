"""
Importación de la cédula EMA/EBA y conciliación trabajador por trabajador (P4-10).

NO HAY LAYOUT OFICIAL ACCESIBLE, Y POR ESO EL PARSER ES CONFIGURABLE
-------------------------------------------------------------------
La EMA y la EBA se descargan de IDSE/Escritorio Virtual como PDF o como el archivo de emisión
que se carga al SUA. # POR CONFIRMAR: no se encontró publicada la estructura de ese archivo
(desde este entorno no hay red hacia imss.gob.mx y ninguna fuente secundaria la transcribe).
En lugar de adivinar posiciones, el parser lee CSV, XLSX o TXT (delimitado o de posiciones
fijas) con una CONFIGURACIÓN que dice qué columna es qué. La plantilla por default
(`PLANTILLAS`) es la que el despacho llena exportando la cédula del SUA a Excel o copiándola
del PDF; si el archivo trae otros encabezados, se mandan en `columnas`.

QUÉ SE COMPARA
--------------
Por NSS: días, SBC e importe total, más cada ramo que la cédula traiga en sus columnas. La
tolerancia de importe es configurable (default un centavo: el motor redondea por concepto y
por trabajador como el SUA, §D2). Lo que sobra o falta de cada lado también es diferencia.
"""

from __future__ import annotations

import csv
import io
import unicodedata
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from typing import Any

from app.exceptions import FiscalAgentError

CERO = Decimal("0.00")

# Columnas de ramo de cada cédula → (ramo del motor, parte).
RAMOS_EMA = {
    "cuota_fija": ("eym_cuota_fija", "patron"),
    "excedente_patronal": ("eym_excedente", "patron"),
    "excedente_obrero": ("eym_excedente", "obrero"),
    "prestaciones_dinero_patronal": ("eym_prestaciones_dinero", "patron"),
    "prestaciones_dinero_obrero": ("eym_prestaciones_dinero", "obrero"),
    "gastos_medicos_patronal": ("eym_gastos_medicos_pensionados", "patron"),
    "gastos_medicos_obrero": ("eym_gastos_medicos_pensionados", "obrero"),
    "riesgos_trabajo": ("riesgos_trabajo", "patron"),
    "invalidez_vida_patronal": ("invalidez_vida", "patron"),
    "invalidez_vida_obrero": ("invalidez_vida", "obrero"),
    "guarderias": ("guarderias", "patron"),
}
RAMOS_EBA = {
    "retiro": ("retiro", "patron"),
    "ceav_patronal": ("ceav", "patron"),
    "ceav_obrero": ("ceav", "obrero"),
    "infonavit_patronal": ("infonavit", "patron"),
}
COMUNES = ("nss", "nombre", "dias", "sbc", "total")
PLANTILLAS = {
    "EMA": (*COMUNES[:4], *RAMOS_EMA, "total"),
    "EBA": (*COMUNES[:4], *RAMOS_EBA, "amortizacion", "total"),
}
NUMERICOS = {"dias", "sbc", "total", "amortizacion", *RAMOS_EMA, *RAMOS_EBA}


@dataclass
class Configuracion:
    """Cómo leer el archivo. Ver el docstring del módulo."""

    formato: str = ""  # csv | xlsx | txt ; vacío = por extensión
    delimitador: str = ","
    fila_encabezado: int = 1  # 1-indexada; 0 = sin encabezado (columnas por número)
    columnas: dict[str, str | int] = field(default_factory=dict)
    posiciones: dict[str, list[int]] = field(default_factory=dict)  # txt de ancho fijo
    decimales_implicitos: int = 0  # p. ej. 2: "000012345" → 123.45
    tolerancia: Decimal = Decimal("0.01")

    @classmethod
    def de_dict(cls, d: dict[str, Any] | None) -> Configuracion:
        d = dict(d or {})
        if "tolerancia" in d:
            d["tolerancia"] = Decimal(str(d["tolerancia"]))
        conocidos = {k: v for k, v in d.items() if k in cls.__dataclass_fields__}
        return cls(**conocidos)

    def como_dict(self) -> dict[str, Any]:
        return {"formato": self.formato, "delimitador": self.delimitador,
                "fila_encabezado": self.fila_encabezado, "columnas": self.columnas,
                "posiciones": self.posiciones, "decimales_implicitos": self.decimales_implicitos,
                "tolerancia": str(self.tolerancia)}


def _normal(texto: Any) -> str:
    s = unicodedata.normalize("NFKD", str(texto or "")).encode("ascii", "ignore").decode()
    return "".join(ch if ch.isalnum() else "_" for ch in s.strip().lower()).strip("_")


def _numero(valor: Any, implicitos: int = 0) -> Decimal | None:
    if valor in (None, ""):
        return None
    if isinstance(valor, int | float | Decimal):
        return Decimal(str(valor))
    limpio = str(valor).replace("$", "").replace(",", "").replace(" ", "").strip()
    if not limpio:
        return None
    try:
        n = Decimal(limpio)
    except InvalidOperation as e:
        raise FiscalAgentError(f"Valor numérico ilegible en la cédula: {valor!r}.", 422) from e
    return n / (Decimal(10) ** implicitos) if implicitos and "." not in limpio else n


def _renglones(datos: bytes, nombre: str, cfg: Configuracion) -> list[list[Any]]:
    formato = cfg.formato or nombre.rsplit(".", 1)[-1].lower()
    if formato == "xlsx":
        from openpyxl import load_workbook

        hoja = load_workbook(io.BytesIO(datos), read_only=True, data_only=True).active
        return [list(f) for f in hoja.iter_rows(values_only=True)] if hoja else []
    texto = datos.decode("utf-8-sig", errors="replace") if datos[:3] == b"\xef\xbb\xbf" else (
        datos.decode("utf-8") if _es_utf8(datos) else datos.decode("cp1252"))
    if formato == "txt" and cfg.posiciones:
        return [[ln] for ln in texto.splitlines() if ln.strip()]
    if formato not in ("csv", "txt"):
        raise FiscalAgentError(f"Formato de cédula no soportado: {formato!r} (csv, xlsx, txt).",
                               422)
    delim = cfg.delimitador or ("|" if formato == "txt" else ",")
    return [r for r in csv.reader(io.StringIO(texto), delimiter=delim) if any(c.strip() for c in r)]


def _es_utf8(datos: bytes) -> bool:
    try:
        datos.decode("utf-8")
    except UnicodeDecodeError:
        return False
    return True


def leer(datos: bytes, nombre: str, tipo: str, cfg: Configuracion) -> list[dict[str, Any]]:
    """Las filas de la cédula con campos canónicos (nss, dias, sbc, total y ramos)."""
    if tipo not in PLANTILLAS:
        raise FiscalAgentError("El tipo de cédula es EMA o EBA.", 422)
    renglones = _renglones(datos, nombre, cfg)
    filas: list[dict[str, Any]] = []
    if cfg.posiciones:
        for (linea,) in renglones:
            crudo = {k: linea[p[0] - 1:p[1]].strip() for k, p in cfg.posiciones.items()}
            filas.append(crudo)
    else:
        if cfg.fila_encabezado:
            if len(renglones) < cfg.fila_encabezado:
                raise FiscalAgentError("La cédula no trae la fila de encabezado indicada.", 422)
            encabezado = [_normal(c) for c in renglones[cfg.fila_encabezado - 1]]
            cuerpo = renglones[cfg.fila_encabezado:]
        else:
            encabezado, cuerpo = [], renglones
        mapa = {campo: campo for campo in PLANTILLAS[tipo]}
        mapa.update(cfg.columnas)
        indices: dict[str, int] = {}
        for campo, col in mapa.items():
            if isinstance(col, int) or (isinstance(col, str) and col.isdigit()):
                indices[campo] = int(col) - 1
            elif _normal(col) in encabezado:
                indices[campo] = encabezado.index(_normal(col))
        if "nss" not in indices:
            raise FiscalAgentError(
                "No encontré la columna del NSS en la cédula. Indica en `columnas` cómo se "
                "llama (p. ej. {\"nss\": \"Número de seguridad social\"}).", 422)
        for r in cuerpo:
            filas.append({c: (r[i] if i < len(r) else None) for c, i in indices.items()})
    salida = []
    for crudo in filas:
        nss = "".join(ch for ch in str(crudo.get("nss") or "") if ch.isdigit())
        if len(nss) != 11:
            continue  # subtotales, encabezados repetidos, renglones de firma
        fila: dict[str, Any] = {"nss": nss, "nombre": str(crudo.get("nombre") or "").strip()}
        for campo in NUMERICOS:
            if campo in crudo:
                n = _numero(crudo[campo], 0 if campo == "dias" else cfg.decimales_implicitos)
                if n is not None:
                    fila[campo] = n
        salida.append(fila)
    if not salida:
        raise FiscalAgentError("La cédula no trae ningún renglón con un NSS de 11 dígitos.", 422)
    return salida


def _ramo_calculado(fila: dict[str, Any], clave: str, parte: str) -> Decimal:
    ramo = fila["ramos"].get(clave)
    return Decimal(ramo[parte]) if ramo else CERO


def conciliar(
    calculo: dict[str, Any], cedula: list[dict[str, Any]], tipo: str, tolerancia: Decimal
) -> dict[str, Any]:
    """Diferencias por trabajador entre el cálculo y la cédula importada."""
    ramos = RAMOS_EMA if tipo == "EMA" else RAMOS_EBA
    nuestros = {f["nss"]: f for f in calculo["trabajadores"] if f.get("nss")}
    sin_nss = [f["empleado_no"] for f in calculo["trabajadores"] if not f.get("nss")]
    suyos = {f["nss"]: f for f in cedula}
    renglones, con_diferencia = [], 0
    for nss in sorted(set(nuestros) | set(suyos)):
        a, b = nuestros.get(nss), suyos.get(nss)
        difs: list[dict[str, str]] = []
        if a is None:
            difs.append({"campo": "trabajador", "calculado": "—", "cedula": "presente",
                         "nota": "Está en la cédula y no en el cálculo (¿baja no presentada o "
                                 "alta que no se capturó?)."})
        elif b is None:
            difs.append({"campo": "trabajador", "calculado": "presente", "cedula": "—",
                         "nota": "Está en el cálculo y no en la cédula (¿alta no reconocida "
                                 "por el IMSS?)."})
        else:
            if "dias" in b and int(b["dias"]) != int(a["dias"]):
                difs.append({"campo": "dias", "calculado": str(a["dias"]),
                             "cedula": str(int(b["dias"]))})
            if "sbc" in b and abs(b["sbc"] - Decimal(a["sbc"])) > Decimal("0.01"):
                difs.append({"campo": "sbc", "calculado": a["sbc"], "cedula": str(b["sbc"])})
            for col, (clave, parte) in ramos.items():
                if col in b:
                    nuestro = _ramo_calculado(a, clave, parte)
                    if abs(b[col] - nuestro) > tolerancia:
                        difs.append({"campo": col, "calculado": str(nuestro),
                                     "cedula": str(b[col]),
                                     "diferencia": str(b[col] - nuestro)})
            if "total" in b:
                total_cedula = b["total"] - b.get("amortizacion", CERO)
                if abs(total_cedula - Decimal(a["total"])) > tolerancia:
                    difs.append({"campo": "importe", "calculado": a["total"],
                                 "cedula": str(total_cedula),
                                 "diferencia": str(total_cedula - Decimal(a["total"]))})
        con_diferencia += bool(difs)
        renglones.append({
            "nss": nss, "empleado_no": a["empleado_no"] if a else "",
            "nombre": (a or {}).get("nombre") or (b or {}).get("nombre", ""),
            "dias_calculo": a["dias"] if a else None,
            "dias_cedula": int(b["dias"]) if b and "dias" in b else None,
            "sbc_calculo": a["sbc"] if a else None,
            "sbc_cedula": str(b["sbc"]) if b and "sbc" in b else None,
            "importe_calculo": a["total"] if a else None,
            "importe_cedula": str(b["total"] - b.get("amortizacion", CERO))
            if b and "total" in b else None,
            "cuadra": not difs, "diferencias": difs,
        })
    total_cedula = sum((f.get("total", CERO) - f.get("amortizacion", CERO) for f in cedula),
                       start=CERO)
    return {"tipo": tipo, "trabajadores": renglones, "con_diferencia": con_diferencia,
            "total_calculo": calculo["total"], "total_cedula": str(total_cedula),
            "diferencia_total": str(total_cedula - Decimal(calculo["total"])),
            "sin_nss_en_ficha": sin_nss, "tolerancia": str(tolerancia)}


def plantilla_csv(tipo: str) -> str:
    """El encabezado de la plantilla por default, para que el despacho la llene."""
    if tipo not in PLANTILLAS:
        raise FiscalAgentError("El tipo de cédula es EMA o EBA.", 422)
    return ",".join(PLANTILLAS[tipo]) + "\r\n"


def reporte_filas(conciliacion: dict[str, Any]) -> list[list[str]]:
    """Renglones planos (uno por diferencia) para CSV/XLSX."""
    salida = [["NSS", "Empleado", "Nombre", "Cuadra", "Campo", "Calculado", "Cédula",
               "Diferencia", "Nota"]]
    for t in conciliacion["trabajadores"]:
        if t["cuadra"]:
            salida.append([t["nss"], t["empleado_no"], t["nombre"], "sí", "", "", "", "", ""])
        for d in t["diferencias"]:
            salida.append([t["nss"], t["empleado_no"], t["nombre"], "no", d["campo"],
                           d.get("calculado", ""), d.get("cedula", ""), d.get("diferencia", ""),
                           d.get("nota", "")])
    salida.append(["TOTAL", "", "", "", "importe", conciliacion["total_calculo"],
                   conciliacion["total_cedula"], conciliacion["diferencia_total"], ""])
    return salida


def a_csv(filas: list[list[str]]) -> bytes:
    buf = io.StringIO()
    csv.writer(buf, lineterminator="\r\n").writerows(filas)
    return buf.getvalue().encode("utf-8-sig")


def a_xlsx(filas: list[list[str]], titulo: str) -> bytes:
    from openpyxl import Workbook

    libro = Workbook()
    hoja = libro.active
    assert hoja is not None
    hoja.title = titulo[:31]
    for f in filas:
        hoja.append(f)
    buf = io.BytesIO()
    libro.save(buf)
    return buf.getvalue()
