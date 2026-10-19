"""
Lectura de estados de cuenta (P2-06): CSV, XLSX y OFX → `MovimientoLeido` normalizado.

Los PDF viven en `pdf_bancos.py`. Todo lector devuelve una `LecturaEstado` con los
movimientos en el orden del archivo, los saldos inicial/final que pudo leer y advertencias
por renglón que no entendió (un renglón ilegible NO tumba la carga: se reporta).

MAPEO DE COLUMNAS (CSV / XLSX)
------------------------------
Cada banco exporta distinto, así que el mapeo es configurable y se guarda como plantilla:

    {"fecha": "Fecha", "concepto": "Descripción", "referencia": "Referencia",
     "cargo": "Retiros", "abono": "Depósitos", "saldo": "Saldo",
     "importe": null, "formato_fecha": "%d/%m/%Y", "fila_encabezado": 1,
     "separador": ",", "signo_importe": "abono_positivo"}

- Una columna se nombra por su encabezado (sin distinguir mayúsculas ni acentos), por su
  letra de Excel ("C") o por su número empezando en 1.
- `importe` (una sola columna con signo) sustituye a `cargo`/`abono`. `signo_importe`
  dice qué significa el positivo: `abono_positivo` (default: + entra) o `cargo_positivo`.
- Sin mapeo se intenta reconocer los encabezados (`detectar_mapeo`).
- `formato_fecha` vacío = se prueban los formatos comunes (dd/mm/aaaa, aaaa-mm-dd,
  dd-MMM-aaaa con mes en español o inglés).
"""

from __future__ import annotations

import csv
import io
import re
import unicodedata
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from typing import Any

from app.exceptions import FiscalAgentError

CERO = Decimal("0")
MAX_BYTES = 10 * 1024 * 1024
MAX_RENGLONES = 20000

MESES = {
    "ene": 1, "feb": 2, "mar": 3, "abr": 4, "may": 5, "jun": 6, "jul": 7, "ago": 8,
    "sep": 9, "set": 9, "oct": 10, "nov": 11, "dic": 12,
    "jan": 1, "apr": 4, "aug": 8, "dec": 12,
}

CAMPOS = ("fecha", "concepto", "referencia", "cargo", "abono", "importe", "saldo")

# Encabezados comunes de los bancos mexicanos → campo (sin acentos, minúsculas).
SINONIMOS: dict[str, tuple[str, ...]] = {
    "fecha": ("fecha", "fecha operacion", "fecha de operacion", "fecha movimiento", "dia",
              "fecha valor", "date"),
    "concepto": ("concepto", "descripcion", "concepto / referencia", "detalle", "movimiento",
                 "descripcion del movimiento", "description"),
    "referencia": ("referencia", "ref", "referencia numerica", "folio", "no. referencia",
                   "numero de referencia", "clave de rastreo", "reference"),
    "cargo": ("cargo", "cargos", "retiro", "retiros", "debito", "debitos", "salida",
              "salidas", "debit"),
    "abono": ("abono", "abonos", "deposito", "depositos", "credito", "creditos", "entrada",
              "entradas", "credit"),
    "importe": ("importe", "monto", "cantidad", "amount"),
    "saldo": ("saldo", "saldo final", "saldo disponible", "balance"),
}


@dataclass
class MovimientoLeido:
    fecha: date
    concepto: str
    referencia: str
    cargo: Decimal
    abono: Decimal
    saldo: Decimal | None = None


@dataclass
class LecturaEstado:
    movimientos: list[MovimientoLeido] = field(default_factory=list)
    saldo_inicial: Decimal | None = None
    saldo_final: Decimal | None = None
    advertencias: list[str] = field(default_factory=list)


# ── utilidades ──


def sin_acentos(texto: str) -> str:
    return "".join(c for c in unicodedata.normalize("NFD", texto)
                   if unicodedata.category(c) != "Mn")


def normalizar(texto: Any) -> str:
    return re.sub(r"\s+", " ", sin_acentos(str(texto or "")).strip().lower())


def importe(valor: Any) -> Decimal | None:
    """'$1,234.56', '(1,234.56)', '-1234.5', 1234.5 → Decimal con 2 decimales; '' → None."""
    if valor is None:
        return None
    if isinstance(valor, (int, float, Decimal)) and not isinstance(valor, bool):
        return Decimal(str(valor)).quantize(Decimal("0.01"))
    texto = str(valor).strip()
    if not texto or texto in ("-", "—"):
        return None
    negativo = texto.startswith("(") and texto.endswith(")")
    texto = texto.strip("()").replace("$", "").replace(",", "").replace(" ", "")
    texto = texto.replace("MXN", "").replace("mxn", "")
    if texto.endswith("-"):
        negativo, texto = True, texto[:-1]
    try:
        d = Decimal(texto)
    except InvalidOperation:
        return None
    if negativo:
        d = -abs(d)
    return d.quantize(Decimal("0.01"))


def fecha(valor: Any, formato: str = "", anio_defecto: int | None = None) -> date | None:
    """Fecha en los formatos comunes de los bancos. `anio_defecto` para '15/ENE'."""
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    texto = str(valor or "").strip()
    if not texto:
        return None
    if formato:
        try:
            return datetime.strptime(texto, formato).date()
        except ValueError:
            return None
    t = normalizar(texto).replace(".", "")
    m = re.fullmatch(r"(\d{4})[-/](\d{1,2})[-/](\d{1,2})(?:[ t].*)?", t)
    if m:
        return _armar(int(m.group(1)), int(m.group(2)), int(m.group(3)))
    m = re.fullmatch(r"(\d{1,2})[-/ ](\d{1,2})[-/ ](\d{2,4})", t)
    if m:
        return _armar(_anio(m.group(3)), int(m.group(2)), int(m.group(1)))
    m = re.fullmatch(r"(\d{1,2})[-/ ]([a-z]{3})[a-z]*(?:[-/ ](\d{2,4}))?", t)
    if m and m.group(2) in MESES:
        anio = _anio(m.group(3)) if m.group(3) else anio_defecto
        if anio is None:
            return None
        return _armar(anio, MESES[m.group(2)], int(m.group(1)))
    return None


def _anio(texto: str) -> int:
    n = int(texto)
    return 2000 + n if n < 100 else n


def _armar(anio: int, mes: int, dia: int) -> date | None:
    try:
        return date(anio, mes, dia)
    except ValueError:
        return None


def validar_tamano(datos: bytes) -> None:
    if not datos:
        raise FiscalAgentError("El archivo está vacío.", 422)
    if len(datos) > MAX_BYTES:
        raise FiscalAgentError("El estado de cuenta pasa de 10 MB.", 422)


# ── tabular (CSV / XLSX) ──


def _letra_a_indice(texto: str) -> int | None:
    if re.fullmatch(r"[A-Za-z]{1,2}", texto):
        n = 0
        for c in texto.upper():
            n = n * 26 + (ord(c) - 64)
        return n - 1
    if re.fullmatch(r"\d{1,3}", texto):
        return int(texto) - 1
    return None


def detectar_mapeo(encabezados: list[Any]) -> dict[str, Any]:
    """Reconoce las columnas por su encabezado. Lo que no reconoce queda fuera."""
    norm = [normalizar(h) for h in encabezados]
    mapeo: dict[str, Any] = {}
    for campo, opciones in SINONIMOS.items():
        for i, h in enumerate(norm):
            if h in opciones and i not in mapeo.values():
                mapeo[campo] = i
                break
    if "cargo" in mapeo or "abono" in mapeo:
        mapeo.pop("importe", None)
    return mapeo


def _resolver_columnas(encabezados: list[Any], mapeo: dict[str, Any]) -> dict[str, int]:
    norm = [normalizar(h) for h in encabezados]
    salida: dict[str, int] = {}
    for campo in CAMPOS:
        v = mapeo.get(campo)
        if v is None or v == "":
            continue
        if isinstance(v, int):
            salida[campo] = v
            continue
        texto = str(v).strip()
        if normalizar(texto) in norm:
            salida[campo] = norm.index(normalizar(texto))
            continue
        indice = _letra_a_indice(texto)
        if indice is None:
            raise FiscalAgentError(
                f"La columna '{texto}' del mapeo ({campo}) no está en el archivo.", 422)
        salida[campo] = indice
    if "fecha" not in salida:
        raise FiscalAgentError("El mapeo necesita la columna de la fecha.", 422)
    if not ({"cargo", "abono"} & set(salida) or "importe" in salida):
        raise FiscalAgentError(
            "El mapeo necesita cargo y abono, o una columna de importe con signo.", 422)
    return salida


def _celda(fila: list[Any], i: int | None) -> Any:
    if i is None or i >= len(fila):
        return None
    return fila[i]


def desde_filas(filas: list[list[Any]], mapeo: dict[str, Any] | None) -> LecturaEstado:
    """Normaliza una tabla (lista de filas) con el mapeo dado o detectado."""
    mapeo = dict(mapeo or {})
    fila_enc = int(mapeo.get("fila_encabezado") or 0)
    if fila_enc <= 0:  # buscar el encabezado en las primeras 30 filas
        fila_enc = next((n + 1 for n, f in enumerate(filas[:30])
                         if "fecha" in detectar_mapeo(f)), 1)
    if fila_enc > len(filas):
        raise FiscalAgentError("El archivo no tiene la fila de encabezados indicada.", 422)
    encabezados = filas[fila_enc - 1]
    if not any(mapeo.get(c) not in (None, "") for c in CAMPOS):
        mapeo.update(detectar_mapeo(encabezados))
    cols = _resolver_columnas(encabezados, mapeo)
    formato = str(mapeo.get("formato_fecha") or "")
    cargo_positivo = mapeo.get("signo_importe") == "cargo_positivo"
    lectura = LecturaEstado()
    for n, fila in enumerate(filas[fila_enc:fila_enc + MAX_RENGLONES], start=fila_enc + 1):
        if not any(str(c or "").strip() for c in fila):
            continue
        f = fecha(_celda(fila, cols["fecha"]), formato)
        concepto = str(_celda(fila, cols.get("concepto")) or "").strip()
        if f is None:
            if re.search(r"saldo\s+(inicial|anterior)", normalizar(concepto)
                         + " " + normalizar(_celda(fila, cols["fecha"]))):
                lectura.saldo_inicial = importe(_celda(fila, cols.get("saldo")))
            elif any(str(c or "").strip() for c in fila):
                lectura.advertencias.append(f"Renglón {n}: fecha ilegible, se omitió.")
            continue
        if "importe" in cols:
            valor = importe(_celda(fila, cols["importe"])) or CERO
            if cargo_positivo:
                valor = -valor
            cargo, abono = (CERO, valor) if valor >= 0 else (-valor, CERO)
        else:
            cargo = abs(importe(_celda(fila, cols.get("cargo"))) or CERO)
            abono = abs(importe(_celda(fila, cols.get("abono"))) or CERO)
        if cargo == 0 and abono == 0:
            lectura.advertencias.append(f"Renglón {n}: sin importe, se omitió.")
            continue
        saldo = importe(_celda(fila, cols.get("saldo")))
        lectura.movimientos.append(MovimientoLeido(
            f, concepto[:300], str(_celda(fila, cols.get("referencia")) or "").strip()[:100],
            cargo, abono, saldo))
    if not lectura.movimientos:
        raise FiscalAgentError("No se encontró ningún movimiento con el mapeo dado.", 422)
    ultimo = lectura.movimientos[-1].saldo
    if ultimo is not None:
        lectura.saldo_final = ultimo
    return lectura


def _decodificar(datos: bytes) -> str:
    for cod in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return datos.decode(cod)
        except UnicodeDecodeError:
            continue
    return datos.decode("latin-1", errors="replace")


def leer_csv(datos: bytes, mapeo: dict[str, Any] | None = None) -> LecturaEstado:
    validar_tamano(datos)
    texto = _decodificar(datos)
    separador = str((mapeo or {}).get("separador") or "")
    if not separador:
        try:
            separador = csv.Sniffer().sniff(texto[:4096], delimiters=",;\t|").delimiter
        except csv.Error:
            separador = ","
    filas = [list(f) for f in csv.reader(io.StringIO(texto), delimiter=separador)]
    return desde_filas(filas, mapeo)


def leer_xlsx(datos: bytes, mapeo: dict[str, Any] | None = None) -> LecturaEstado:
    validar_tamano(datos)
    try:
        from openpyxl import load_workbook

        libro = load_workbook(io.BytesIO(datos), read_only=True, data_only=True)
    except Exception as e:  # noqa: BLE001 — cualquier XLSX dañado es un 422 legible
        raise FiscalAgentError("El archivo no es un XLSX válido.", 422) from e
    hoja_nombre = str((mapeo or {}).get("hoja") or "")
    hoja = libro[hoja_nombre] if hoja_nombre in libro.sheetnames else libro.worksheets[0]
    filas = [list(f) for f in hoja.iter_rows(values_only=True, max_row=MAX_RENGLONES + 50)]
    libro.close()
    return desde_filas(filas, mapeo)


# ── OFX ──


def _ofx_valor(bloque: str, etiqueta: str) -> str:
    m = re.search(rf"<{etiqueta}>([^<\r\n]*)", bloque, re.IGNORECASE)
    return m.group(1).strip() if m else ""


def leer_ofx(datos: bytes) -> LecturaEstado:
    """OFX 1.x (SGML) o 2.x (XML): STMTTRN → movimiento; LEDGERBAL → saldo final."""
    validar_tamano(datos)
    texto = _decodificar(datos)
    if "<OFX>" not in texto.upper():
        raise FiscalAgentError("El archivo no es un OFX.", 422)
    lectura = LecturaEstado()
    for bloque in re.findall(r"<STMTTRN>(.*?)(?:</STMTTRN>|(?=<STMTTRN>)|</BANKTRANLIST>)",
                             texto, re.IGNORECASE | re.DOTALL):
        crudo = _ofx_valor(bloque, "DTPOSTED")[:8]
        f = _armar(int(crudo[:4]), int(crudo[4:6]), int(crudo[6:8])) if crudo.isdigit() \
            and len(crudo) == 8 else None
        monto = importe(_ofx_valor(bloque, "TRNAMT"))
        if f is None or monto is None or monto == 0:
            lectura.advertencias.append("Un STMTTRN sin fecha o importe se omitió.")
            continue
        concepto = " ".join(x for x in (_ofx_valor(bloque, "NAME"), _ofx_valor(bloque, "MEMO"))
                            if x)
        referencia = _ofx_valor(bloque, "CHECKNUM") or _ofx_valor(bloque, "REFNUM") \
            or _ofx_valor(bloque, "FITID")
        cargo, abono = (CERO, monto) if monto > 0 else (-monto, CERO)
        lectura.movimientos.append(MovimientoLeido(f, concepto[:300], referencia[:100],
                                                   cargo, abono, None))
    if not lectura.movimientos:
        raise FiscalAgentError("El OFX no trae movimientos.", 422)
    m = re.search(r"<LEDGERBAL>(.*?)(?:</LEDGERBAL>|<AVAILBAL>|$)", texto,
                  re.IGNORECASE | re.DOTALL)
    if m:
        lectura.saldo_final = importe(_ofx_valor(m.group(1), "BALAMT"))
        if lectura.saldo_final is not None:
            flujo = sum((x.abono - x.cargo for x in lectura.movimientos), CERO)
            lectura.saldo_inicial = lectura.saldo_final - flujo
    return lectura
