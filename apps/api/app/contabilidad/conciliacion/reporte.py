"""
Reporte de conciliación bancaria al corte (P2-06).

    Saldo según el banco
      + depósitos registrados en libros que el banco aún no refleja
      − retiros (cheques, transferencias) registrados en libros que el banco aún no refleja
      = saldo del banco conciliado
    Saldo según libros (cuenta contable ligada, con saldos iniciales)
      + abonos del banco no registrados en libros (intereses, depósitos sin identificar)
      − cargos del banco no registrados en libros (comisiones, cargos sin identificar)
      = saldo de libros conciliado
    Diferencia = banco conciliado − libros conciliado (0 = cuadra)

PENDIENTE AL CORTE
------------------
Una partida con fecha ≤ corte está en tránsito si no está conciliada, o si su conciliación
la casó con algo fechado DESPUÉS del corte (un cheque de marzo cobrado en abril sigue en
tránsito en la conciliación de marzo). Un movimiento del banco conciliado sólo contra CFDI
de la bóveda (sin póliza) sigue sin estar en libros: cuenta como "banco no registrado en
libros" con la nota "CFDI sin póliza".

SALDO DEL BANCO
---------------
El saldo del último movimiento ≤ corte que traiga saldo, más el flujo de los que vengan
después sin saldo. Sin ningún saldo en los estados de cuenta, se parte del saldo inicial del
estado más antiguo (o 0) y se advierte.
"""

from __future__ import annotations

from datetime import date
from decimal import Decimal
from typing import Any

from sqlalchemy import select
from sqlalchemy.orm import Session

from app.contabilidad import balanza, polizas
from app.contabilidad.conciliacion.servicio import (
    CERO,
    movimientos_banco,
    movimientos_libros,
)
from app.exportar.excel_seguro import libro_a_bytes
from app.modelos.cfdi import Comprobante
from app.modelos.conciliacion import (
    Conciliacion,
    CuentaBancaria,
    EstadoCuenta,
    MovimientoBancario,
    PartidaConciliacion,
)
from app.modelos.contabilidad import CuentaContable, MovimientoPoliza, Poliza
from app.modelos.nucleo import Cliente


def _fechas_por_conciliacion(db: Session, cb: CuentaBancaria) -> dict[int, list[tuple[str, date]]]:
    """conciliación → [(clave, fecha)] de todas sus partidas."""
    salida: dict[int, list[tuple[str, date]]] = {}
    filas = db.execute(
        select(PartidaConciliacion, MovimientoBancario.fecha, Poliza.fecha, Comprobante.fecha)
        .join(Conciliacion, Conciliacion.id == PartidaConciliacion.conciliacion_id)
        .outerjoin(MovimientoBancario,
                   MovimientoBancario.id == PartidaConciliacion.movimiento_bancario_id)
        .outerjoin(MovimientoPoliza,
                   MovimientoPoliza.id == PartidaConciliacion.movimiento_poliza_id)
        .outerjoin(Poliza, Poliza.id == MovimientoPoliza.poliza_id)
        .outerjoin(Comprobante, Comprobante.id == PartidaConciliacion.comprobante_id)
        .where(Conciliacion.cuenta_bancaria_id == cb.id))
    for p, fb, fp, fc in filas:
        if p.movimiento_bancario_id:
            clave, f = f"b:{p.movimiento_bancario_id}", fb
        elif p.movimiento_poliza_id:
            clave, f = f"p:{p.movimiento_poliza_id}", fp
        else:
            clave, f = f"c:{p.comprobante_id}", fc.date() if fc else None
        salida.setdefault(p.conciliacion_id, []).append((clave, f))
    return salida


def _saldo_banco(db: Session, cb: CuentaBancaria, movs: list[MovimientoBancario],
                 advertencias: list[str]) -> Decimal:
    saldo: Decimal | None = None
    flujo = CERO
    for m in movs:
        if m.saldo is not None:
            saldo, flujo = m.saldo, CERO
        else:
            flujo += m.abono - m.cargo
    if saldo is not None:
        return saldo + flujo
    primero = db.scalar(select(EstadoCuenta).where(EstadoCuenta.cuenta_bancaria_id == cb.id)
                        .order_by(EstadoCuenta.desde, EstadoCuenta.id))
    inicial = primero.saldo_inicial if primero and primero.saldo_inicial is not None else CERO
    if primero is None or primero.saldo_inicial is None:
        advertencias.append("Los estados de cuenta no traen saldo: el saldo del banco parte "
                            "de 0 más los movimientos importados.")
    return inicial + flujo


def reporte(db: Session, cliente: Cliente, cb: CuentaBancaria, corte: date) -> dict[str, Any]:
    advertencias: list[str] = []
    conc_de: dict[str, int] = {}
    fechas = _fechas_por_conciliacion(db, cb)
    for cid, partidas in fechas.items():
        for clave, _f in partidas:
            conc_de[clave] = cid

    def en_transito(clave: str) -> bool:
        cid = conc_de.get(clave)
        if cid is None:
            return True
        return any(f is None or f > corte for _k, f in fechas[cid])

    def solo_cfdi(clave: str) -> bool:
        cid = conc_de.get(clave)
        return cid is not None and all(not k.startswith("p:") for k, _f in fechas[cid]
                                       if not k.startswith("b:"))

    movs_b = movimientos_banco(db, cb, None, corte)
    saldo_banco = _saldo_banco(db, cb, movs_b, advertencias)
    hojas = balanza.totales_por_cuenta(db, cliente, corte, corte)
    t = hojas.get(cb.cuenta_contable_id)
    saldo_libros = t.final if t else CERO

    dep_libros, ret_libros, abo_banco, car_banco = [], [], [], []
    for m, p in movimientos_libros(db, cliente, cb, None, corte):
        if not en_transito(f"p:{m.id}"):
            continue
        fila = {"fecha": p.fecha.isoformat(), "folio": polizas.folio(p),
                "concepto": m.concepto or p.concepto, "importe": abs(m.cargo - m.abono)}
        (dep_libros if m.cargo > 0 else ret_libros).append(fila)
    for m in movs_b:
        clave = f"b:{m.id}"
        if not (en_transito(clave) or solo_cfdi(clave)):
            continue
        fila = {"fecha": m.fecha.isoformat(), "folio": m.referencia,
                "concepto": m.concepto + (" (CFDI sin póliza)" if solo_cfdi(clave) else ""),
                "importe": m.abono if m.abono > 0 else m.cargo}
        (abo_banco if m.abono > 0 else car_banco).append(fila)

    def total(filas: list[dict[str, Any]]) -> Decimal:
        return sum((f["importe"] for f in filas), CERO)

    banco_conc = saldo_banco + total(dep_libros) - total(ret_libros)
    libros_conc = saldo_libros + total(abo_banco) - total(car_banco)
    cuenta = db.get(CuentaContable, cb.cuenta_contable_id)

    def fmt(filas: list[dict[str, Any]]) -> list[dict[str, Any]]:
        return [{**f, "importe": f"{f['importe']:.2f}"} for f in filas]

    return {
        "corte": corte.isoformat(),
        "cuenta_bancaria": {"id": cb.id, "banco": cb.banco, "nombre": cb.nombre,
                            "numero": cb.numero, "cuenta_contable": cuenta.numero if cuenta
                            else ""},
        "saldo_banco": f"{saldo_banco:.2f}", "saldo_libros": f"{saldo_libros:.2f}",
        "depositos_en_transito": fmt(dep_libros),
        "retiros_en_transito": fmt(ret_libros),
        "abonos_banco_no_registrados": fmt(abo_banco),
        "cargos_banco_no_registrados": fmt(car_banco),
        "totales": {"depositos_en_transito": f"{total(dep_libros):.2f}",
                    "retiros_en_transito": f"{total(ret_libros):.2f}",
                    "abonos_banco_no_registrados": f"{total(abo_banco):.2f}",
                    "cargos_banco_no_registrados": f"{total(car_banco):.2f}"},
        "saldo_banco_conciliado": f"{banco_conc:.2f}",
        "saldo_libros_conciliado": f"{libros_conc:.2f}",
        "diferencia": f"{banco_conc - libros_conc:.2f}",
        "cuadra": banco_conc == libros_conc,
        "advertencias": advertencias,
    }


SECCIONES = (
    ("depositos_en_transito", "(+) Depósitos en libros no reflejados por el banco"),
    ("retiros_en_transito", "(−) Retiros en libros no reflejados por el banco"),
    ("abonos_banco_no_registrados", "(+) Abonos del banco no registrados en libros"),
    ("cargos_banco_no_registrados", "(−) Cargos del banco no registrados en libros"),
)


def a_excel(r: dict[str, Any], nombre_cliente: str) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    libro = Workbook()
    hoja = libro.active
    hoja.title = "Conciliación"
    negrita = Font(bold=True)
    cb = r["cuenta_bancaria"]
    hoja.append([f"Conciliación bancaria — {nombre_cliente}"])
    hoja["A1"].font = Font(bold=True, size=13)
    hoja.append([f"Cuenta {cb['nombre']} {cb['numero']} (contable {cb['cuenta_contable']})",
                 "", "", f"Corte {r['corte']}"])
    hoja.append([])

    def renglon(texto: str, valor: str, bold: bool = False) -> None:
        hoja.append([texto, "", "", Decimal(valor)])
        if bold:
            hoja.cell(hoja.max_row, 1).font = negrita
            hoja.cell(hoja.max_row, 4).font = negrita

    renglon("Saldo según el banco", r["saldo_banco"], True)
    for clave, titulo in SECCIONES[:2]:
        _seccion(hoja, titulo, r[clave], r["totales"][clave], negrita)
    renglon("Saldo del banco conciliado", r["saldo_banco_conciliado"], True)
    hoja.append([])
    renglon("Saldo según libros", r["saldo_libros"], True)
    for clave, titulo in SECCIONES[2:]:
        _seccion(hoja, titulo, r[clave], r["totales"][clave], negrita)
    renglon("Saldo de libros conciliado", r["saldo_libros_conciliado"], True)
    hoja.append([])
    renglon("Diferencia", r["diferencia"], True)
    for col, ancho in zip("ABCD", (14, 18, 60, 16), strict=True):
        hoja.column_dimensions[col].width = ancho
    for fila in hoja.iter_rows(min_col=4, max_col=4):
        for celda in fila:
            celda.number_format = "#,##0.00"
    return libro_a_bytes(libro)  # conceptos y referencias del banco: sin fórmulas


def _seccion(hoja: Any, titulo: str, filas: list[dict[str, Any]], total: str,
             negrita: Any) -> None:
    hoja.append([titulo])
    hoja.cell(hoja.max_row, 1).font = negrita
    for f in filas:
        hoja.append([f["fecha"], f["folio"], f["concepto"], Decimal(f["importe"])])
    hoja.append(["", "", "Total", Decimal(total)])
