"""
Exportación a Excel de los reportes contables (P2-05), con openpyxl.

Los importes se escriben como NÚMEROS (no texto) con formato `#,##0.00`, de modo que el
contador puede sumar y filtrar en Excel; los totales se escriben con el valor que calculó el
backend (no con fórmulas), para que el archivo diga exactamente lo que dice la pantalla
aunque alguien lo abra en un visor que no recalcula. Cada hoja lleva encabezado con
cliente, RFC, periodo y moneda.
"""

from __future__ import annotations

import io
from decimal import Decimal
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

FORMATO = "#,##0.00;[Red]-#,##0.00"
NEGRITA = Font(bold=True)
TITULO = Font(bold=True, size=14)
ENCABEZADO = PatternFill("solid", fgColor="374F64")
BLANCO = Font(bold=True, color="FFFFFF")
LINEA = Border(top=Side(style="thin"))
DOBLE = Border(top=Side(style="thin"), bottom=Side(style="double"))
MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _n(v: Any) -> Decimal:
    return Decimal(str(v or "0"))


def _cabecera(ws: Worksheet, titulo: str, cliente: dict[str, Any], periodo: str,
              moneda: str = "MXN") -> int:
    ws["A1"] = titulo
    ws["A1"].font = TITULO
    ws["A2"] = f"{cliente.get('nombre', '')} · RFC {cliente.get('rfc', '')}"
    ws["A3"] = f"{periodo} · Moneda: {moneda}"
    return 5


def _encabezados(ws: Worksheet, fila: int, titulos: list[str]) -> None:
    for i, t in enumerate(titulos, start=1):
        celda = ws.cell(row=fila, column=i, value=t)
        celda.font, celda.fill = BLANCO, ENCABEZADO
        celda.alignment = Alignment(horizontal="center" if i > 2 else "left", wrap_text=True)


def _importes(ws: Worksheet, fila: int, col: int, valores: list[Any], *, negrita: bool = False,
              borde: Border | None = None) -> None:
    for i, v in enumerate(valores):
        celda = ws.cell(row=fila, column=col + i, value=float(_n(v)))
        celda.number_format = FORMATO
        if negrita:
            celda.font = NEGRITA
        if borde:
            celda.border = borde


def _anchos(ws: Worksheet, anchos: list[int]) -> None:
    for i, a in enumerate(anchos, start=1):
        ws.column_dimensions[get_column_letter(i)].width = a


def _bytes(wb: Workbook) -> bytes:
    salida = io.BytesIO()
    wb.save(salida)
    return salida.getvalue()


def balanza(datos: dict[str, Any], cliente: dict[str, Any]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Balanza"
    fila = _cabecera(ws, "Balanza de comprobación", cliente,
                     f"Del {datos['desde']} al {datos['hasta']}")
    cols = ["inicial_deudor", "inicial_acreedor", "cargos", "abonos", "final_deudor",
            "final_acreedor"]
    _encabezados(ws, fila, ["Cuenta", "Nombre", "Saldo inicial deudor",
                            "Saldo inicial acreedor", "Cargos", "Abonos",
                            "Saldo final deudor", "Saldo final acreedor"])
    ws.freeze_panes = ws.cell(row=fila + 1, column=3)
    for r in datos["renglones"]:
        fila += 1
        ws.cell(row=fila, column=1, value=r["numero"])
        ws.cell(row=fila, column=2, value=r["nombre"]).alignment = Alignment(
            indent=max(0, r["nivel"] - 1))
        _importes(ws, fila, 3, [r[c] for c in cols], negrita=r["nivel"] == 1)
        if r["nivel"] == 1:
            ws.cell(row=fila, column=1).font = ws.cell(row=fila, column=2).font = NEGRITA
    fila += 1
    ws.cell(row=fila, column=1, value="Totales").font = NEGRITA
    _importes(ws, fila, 3, [datos["totales"][c] for c in cols], negrita=True, borde=DOBLE)
    ws.cell(row=fila + 2, column=1,
            value="Cuadra" if datos["cuadra"] else "NO CUADRA").font = NEGRITA
    _anchos(ws, [16, 44, 18, 18, 18, 18, 18, 18])
    return _bytes(wb)


def auxiliar(datos: dict[str, Any], cliente: dict[str, Any]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Auxiliar"
    cuenta = datos["cuenta"]
    fila = _cabecera(ws, f"Auxiliar {cuenta['numero']} {cuenta['nombre']}", cliente,
                     f"Del {datos['desde']} al {datos['hasta']}")
    _encabezados(ws, fila, ["Fecha", "Póliza", "Cuenta", "Concepto", "UUID", "Cargo", "Abono",
                            "Saldo"])
    fila += 1
    ws.cell(row=fila, column=4, value="Saldo inicial").font = NEGRITA
    _importes(ws, fila, 8, [datos["saldo_inicial"]], negrita=True)
    for m in datos["movimientos"]:
        fila += 1
        for i, v in enumerate([m["fecha"], m["folio"], m["cuenta"], m["concepto"], m["uuid"]],
                              start=1):
            ws.cell(row=fila, column=i, value=v)
        _importes(ws, fila, 6, [m["cargo"], m["abono"], m["saldo"]])
    fila += 1
    ws.cell(row=fila, column=4, value="Totales y saldo final").font = NEGRITA
    _importes(ws, fila, 6, [datos["cargos"], datos["abonos"], datos["saldo_final"]],
              negrita=True, borde=DOBLE)
    _anchos(ws, [12, 18, 14, 44, 38, 16, 16, 16])
    return _bytes(wb)


def estado_financiero(datos: dict[str, Any]) -> bytes:
    """Estado de resultados o balance general (las líneas que arma `reportes.py`)."""
    wb = Workbook()
    ws = wb.active
    ws.title = datos["titulo"][:31]
    periodo = f"Del {datos['desde']} al {datos['hasta']}"
    if datos["reporte"] == "balance_general":
        periodo = f"Al {datos['hasta']}"
    if datos.get("comparativo"):
        c = datos["comparativo"]
        periodo += f" · comparativo {c['desde']} a {c['hasta']}"
    fila = _cabecera(ws, datos["titulo"], datos["cliente"], periodo, datos["moneda"])
    if datos.get("tipo_cambio"):
        ws["A4"] = f"Tipo de cambio: {datos['tipo_cambio']} MXN por 1 {datos['moneda']}"
    columnas = datos["columnas"]
    _encabezados(ws, fila, ["Cuenta", "Concepto", *[c["etiqueta"] for c in columnas]])
    for linea in datos["lineas"]:
        fila += 1
        tipo = linea["tipo"]
        if tipo == "encabezado":
            ws.cell(row=fila, column=1, value=linea["etiqueta"]).font = NEGRITA
            continue
        valores = [linea["valores"][c["clave"]] for c in columnas]
        if tipo == "cuenta":
            ws.cell(row=fila, column=1, value=linea["numero"])
            ws.cell(row=fila, column=2, value=linea["etiqueta"]).alignment = Alignment(
                indent=max(0, linea["nivel"] - 1))
            _importes(ws, fila, 3, valores)
            continue
        ws.cell(row=fila, column=2, value=linea["etiqueta"]).font = NEGRITA
        _importes(ws, fila, 3, valores, negrita=True, borde=DOBLE if tipo == "total" else LINEA)
    if datos["reporte"] == "balance_general":
        fila += 2
        ws.cell(row=fila, column=2, value="Activo = pasivo + capital: " +
                ("cuadra" if datos["cuadra"] else f"diferencia {datos['diferencia']}")
                ).font = NEGRITA
    _anchos(ws, [16, 48, *[20 for _ in columnas]])
    return _bytes(wb)


def polizas(lista: list[dict[str, Any]], cliente: dict[str, Any], desde: str,
            hasta: str) -> bytes:
    """Pólizas con sus movimientos (impresión de pólizas del periodo)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Pólizas"
    fila = _cabecera(ws, "Pólizas", cliente, f"Del {desde} al {hasta}")
    _encabezados(ws, fila, ["Folio", "Fecha", "Cuenta", "Nombre de la cuenta", "Concepto",
                            "UUID", "Cargo", "Abono"])
    for p in lista:
        fila += 1
        ws.cell(row=fila, column=1, value=p["folio"]).font = NEGRITA
        ws.cell(row=fila, column=2, value=p["fecha"])
        estado = "" if p["estado"] == "activa" else " (CANCELADA)"
        ws.cell(row=fila, column=5, value=f"{p['concepto']}{estado}").font = NEGRITA
        for m in p.get("movimientos", []):
            fila += 1
            for i, v in enumerate([m["cuenta"], m["nombre_cuenta"], m["concepto"], m["uuid"]],
                                  start=3):
                ws.cell(row=fila, column=i, value=v)
            _importes(ws, fila, 7, [m["cargo"], m["abono"]])
        fila += 1
        ws.cell(row=fila, column=5, value="Total póliza").font = NEGRITA
        _importes(ws, fila, 7, [p["total_cargos"], p["total_abonos"]], negrita=True,
                  borde=LINEA)
    _anchos(ws, [18, 12, 14, 34, 40, 38, 16, 16])
    return _bytes(wb)
