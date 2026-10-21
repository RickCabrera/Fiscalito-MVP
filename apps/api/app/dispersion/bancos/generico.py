"""
Layout genérico CSV / XLSX configurable por columnas (P4-12): el respaldo de todos.

Sirve para el banco que no está en el catálogo, para los portales que aceptan un CSV propio
(y se configuran por columnas), y como hoja de control para el contador. Es formato propio:
no hay autoridad que lo defina y por eso se puede afirmar entero (`verificado`).

El importe va con punto decimal y dos decimales (`1234.56`), o en centavos si se elige la
columna `importe_centavos`. Al menos una de las dos es obligatoria: sin importe no hay cuadre.
"""

from __future__ import annotations

import csv
import io
from decimal import Decimal

from openpyxl import Workbook, load_workbook

from app.dispersion.clabe import nombre_banco
from app.dispersion.tipos import Abono, Cargo, Layout, Opciones, Pago
from app.exceptions import FiscalAgentError

COLUMNAS: dict[str, str] = {
    "empleado_no": "Número de empleado",
    "nombre": "Nombre completo",
    "apellido_paterno": "Apellido paterno",
    "apellido_materno": "Apellido materno",
    "nombres": "Nombre(s)",
    "rfc": "RFC",
    "banco": "Clave del banco (c_Banco)",
    "banco_nombre": "Banco",
    "tipo_cuenta": "Tipo de cuenta (cuenta, tarjeta, clabe)",
    "cuenta": "Cuenta de abono",
    "clabe": "CLABE",
    "numero_cuenta": "Número de cuenta o tarjeta",
    "importe": "Importe (con punto decimal)",
    "importe_centavos": "Importe en centavos",
    "concepto": "Concepto",
    "fecha_aplicacion": "Fecha de aplicación",
}
DEFAULT = ("empleado_no", "nombre", "banco_nombre", "clabe", "numero_cuenta", "importe",
           "concepto")
SEPARADORES = {",": ",", "|": "|", ";": ";", "tab": "\t"}


def _importe(centavos: int) -> str:
    return f"{centavos // 100}.{centavos % 100:02d}"


class Generico(Layout):
    id = "generico"
    nombre = "Genérico"
    banco = None
    estado = "verificado"
    fuente = "Formato propio configurable por columnas; sus importes son los netos guardados."
    interbancario = True

    def __init__(self, columnas: list[str] | None = None, separador: str = ",",
                 formato: str = "csv", encabezado: bool = True) -> None:
        self.columnas = list(columnas or DEFAULT)
        desconocidas = [c for c in self.columnas if c not in COLUMNAS]
        if desconocidas or not self.columnas:
            raise FiscalAgentError(f"Columnas que no existen: {desconocidas}.", 422)
        if "importe" not in self.columnas and "importe_centavos" not in self.columnas:
            raise FiscalAgentError(
                "El layout genérico necesita la columna 'importe' o 'importe_centavos': sin "
                "importe no hay cuadre contra la nómina.", 422)
        if separador not in SEPARADORES:
            raise FiscalAgentError(f"Separador no admitido: {separador!r}.", 422)
        if formato not in ("csv", "xlsx"):
            raise FiscalAgentError("El formato genérico es csv o xlsx.", 422)
        self.separador, self.formato, self.encabezado = SEPARADORES[separador], formato, encabezado
        self.extension = formato  # type: ignore[misc]
        self.mime = (  # type: ignore[misc]
            "text/csv" if formato == "csv"
            else "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

    def abono(self, pago: Pago) -> Abono | str:
        if pago.clabe:
            return Abono("clabe", pago.clabe, pago.banco, False)
        if pago.numero_cuenta:
            tipo = "tarjeta" if len(pago.numero_cuenta) == 16 else "cuenta"
            return Abono(tipo, pago.numero_cuenta, pago.banco, False)
        return "Sin cuenta ni CLABE."

    def _fila(self, p: Pago, a: Abono, op: Opciones) -> list[str]:
        valores = {
            "empleado_no": p.empleado_no, "nombre": p.nombre,
            "apellido_paterno": p.apellido_paterno, "apellido_materno": p.apellido_materno,
            "nombres": p.nombres, "rfc": p.rfc, "banco": a.banco,
            "banco_nombre": nombre_banco(a.banco), "tipo_cuenta": a.tipo, "cuenta": a.numero,
            "clabe": p.clabe, "numero_cuenta": p.numero_cuenta, "importe": _importe(p.centavos),
            "importe_centavos": str(p.centavos), "concepto": op.concepto,
            "fecha_aplicacion": op.fecha_aplicacion.isoformat(),
        }
        return [valores[c] for c in self.columnas]

    def generar(self, pagos: list[tuple[Pago, Abono]], cargo: Cargo, op: Opciones) -> bytes:
        filas = [self._fila(p, a, op) for p, a in pagos]
        titulos = [COLUMNAS[c] for c in self.columnas]
        if self.formato == "csv":
            buf = io.StringIO()
            w = csv.writer(buf, delimiter=self.separador, lineterminator="\r\n")
            if self.encabezado:
                w.writerow(titulos)
            w.writerows(filas)
            # UTF-8 con BOM: Excel en español abre bien los acentos así.
            return buf.getvalue().encode("utf-8-sig")
        libro = Workbook()
        hoja = libro.active
        hoja.title = "Dispersion"
        if self.encabezado:
            hoja.append(titulos)
        for fila in filas:
            # Cuentas y CLABE como texto: como número, Excel les come los ceros a la izquierda.
            hoja.append([Decimal(v) if c == "importe" else v
                         for c, v in zip(self.columnas, fila, strict=True)])
        if "importe" in self.columnas:
            col = self.columnas.index("importe") + 1
            for (celda,) in hoja.iter_rows(min_row=2 if self.encabezado else 1,
                                           min_col=col, max_col=col):
                celda.number_format = "0.00"
        salida = io.BytesIO()
        libro.save(salida)
        return salida.getvalue()

    def importes(self, contenido: bytes) -> list[int]:
        if self.formato == "csv":
            filas = list(csv.reader(io.StringIO(contenido.decode("utf-8-sig")),
                                    delimiter=self.separador))
        else:
            hoja = load_workbook(io.BytesIO(contenido)).active
            filas = [[("" if v is None else str(v)) for v in r]
                     for r in hoja.iter_rows(values_only=True)]
        if self.encabezado:
            filas = filas[1:]
        if "importe_centavos" in self.columnas:
            i = self.columnas.index("importe_centavos")
            return [int(f[i]) for f in filas]
        i = self.columnas.index("importe")
        return [int((Decimal(f[i]) * 100).to_integral_exact()) for f in filas]
